# -*- coding: utf-8 -*-
import os
import sys
import io

try:
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def apply_document_page_setup(doc):
    for sec in doc.sections:
        sec.top_margin = Inches(0.75)
        sec.bottom_margin = Inches(0.75)
        sec.left_margin = Inches(0.75)
        sec.right_margin = Inches(0.75)
        sec.page_width = Inches(8.27)   # A4
        sec.page_height = Inches(11.69) # A4

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="94A3B8", sz="6", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_header_info_table(doc):
    t = doc.add_table(rows=8, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t, color="94A3B8", sz="6")
    
    info_data = [
        ("Sinh viên thực hiện", "Trần Văn Hoàng"),
        ("Mã số sinh viên (MSSV)", "0034467"),
        ("Lớp chuyên ngành", "67CS – Khóa K67"),
        ("Đơn vị thực tập", "Công ty TNHH Nền tảng chuyển đổi số Việt Nam"),
        ("Bộ phận công tác", "eGov – Khối Chính phủ số"),
        ("Cán bộ hướng dẫn tại đơn vị", "Đào Đức Văn"),
        ("Dự án tham gia", "Hệ thống báo cáo điện tử eForm"),
        ("Đề tài nghiên cứu", "Nghiên cứu và ứng dụng AI xây dựng Module Import file Excel thông minh, tự động ánh xạ dữ liệu cho hệ thống Báo cáo thống kê eForm.")
    ]
    
    col_widths = [Inches(2.1), Inches(4.6)]
    for i, (label, val) in enumerate(info_data):
        row = t.rows[i]
        row.cells[0].width = col_widths[0]
        row.cells[1].width = col_widths[1]
        
        set_cell_background(row.cells[0], "F8FAFC")
        set_cell_margins(row.cells[0], top=70, bottom=70, left=120, right=120)
        set_cell_margins(row.cells[1], top=70, bottom=70, left=120, right=120)
        
        p0 = row.cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r0 = p0.add_run(label)
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(11.5)
        r0.font.bold = True
        r0.font.color.rgb = RGBColor(30, 41, 59)
        
        p1 = row.cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1 = p1.add_run(val)
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(11.5)
        if label in ["Sinh viên thực hiện", "Mã số sinh viên (MSSV)", "Cán bộ hướng dẫn tại đơn vị"]:
            r1.font.bold = True
        r1.font.color.rgb = RGBColor(15, 23, 42)
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(2)
    p_sp.paragraph_format.space_after = Pt(6)

def add_title(doc, week_num, date_range, subtitle):
    p0 = doc.add_paragraph()
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = p0.add_run(f"BÁO CÁO THỰC TẬP TUẦN {week_num}")
    r0.font.name = "Times New Roman"
    r0.font.size = Pt(18)
    r0.font.bold = True
    r0.font.color.rgb = RGBColor(15, 23, 42)
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(2)
    
    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_date = p_date.add_run(f"({date_range})")
    r_date.font.name = "Times New Roman"
    r_date.font.size = Pt(11.5)
    r_date.font.italic = True
    r_date.font.color.rgb = RGBColor(100, 116, 139)
    p_date.paragraph_format.space_before = Pt(0)
    p_date.paragraph_format.space_after = Pt(4)
    
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run(subtitle.upper())
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(13)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(2, 132, 199)
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(14)

def add_h1(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(5)
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 41, 59)
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(3)
    return p

def add_h3(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.italic = True
    r.font.color.rgb = RGBColor(51, 65, 85)
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    return p

def add_p(doc, text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    
    if bold_prefix:
        rb = p.add_run(bold_prefix + " ")
        rb.font.name = "Times New Roman"
        rb.font.size = Pt(12)
        rb.font.bold = True
        rb.font.color.rgb = RGBColor(15, 23, 42)
        
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.italic = italic
    r.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_bullet(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.2
    
    if bold_prefix:
        rb = p.add_run(bold_prefix + ": ")
        rb.font.name = "Times New Roman"
        rb.font.size = Pt(12)
        rb.font.bold = True
        rb.font.color.rgb = RGBColor(15, 23, 42)
        
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_image_figure(doc, img_path, caption):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(3)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(6.2))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(8)
        r_cap = p_cap.add_run(caption)
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(10.5)
        r_cap.font.italic = True
        r_cap.font.bold = True
        r_cap.font.color.rgb = RGBColor(51, 65, 85)
    else:
        print(f"Lưu ý: Không tìm thấy ảnh {img_path}")

def add_custom_table(doc, headers, rows_data, col_widths=None):
    t = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t, color="94A3B8", sz="6")
    
    # Header
    hdr_row = t.rows[0]
    for j, h in enumerate(headers):
        if col_widths:
            hdr_row.cells[j].width = col_widths[j]
        set_cell_background(hdr_row.cells[j], "E2E8F0")
        set_cell_margins(hdr_row.cells[j], top=70, bottom=70, left=90, right=90)
        p = hdr_row.cells[j].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = RGBColor(15, 23, 42)
        
    # Rows
    for i, row_data in enumerate(rows_data):
        row = t.rows[i + 1]
        bg_col = "FFFFFF" if i % 2 == 0 else "F8FAFC"
        for j, val in enumerate(row_data):
            if col_widths:
                row.cells[j].width = col_widths[j]
            set_cell_background(row.cells[j], bg_col)
            set_cell_margins(row.cells[j], top=55, bottom=55, left=90, right=90)
            p = row.cells[j].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j > 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(val))
            r.font.name = "Times New Roman"
            r.font.size = Pt(10.5)
            if j == len(row_data) - 1 and str(val) in ["Đạt", "Hoàn thành", "Tốt", "Chính xác"]:
                r.font.bold = True
                r.font.color.rgb = RGBColor(16, 185, 129)
            elif j == len(row_data) - 1 and str(val) in ["Chưa đạt", "Lỗi"]:
                r.font.bold = True
                r.font.color.rgb = RGBColor(239, 68, 68)
            elif j == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(30, 41, 59)
            else:
                r.font.color.rgb = RGBColor(51, 65, 85)
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(2)
    p_sp.paragraph_format.space_after = Pt(6)

def generate_week_4():
    doc = Document()
    apply_document_page_setup(doc)
    
    add_title(
        doc,
        week_num=4,
        date_range="Từ ngày 24/08/2026 đến ngày 28/08/2026",
        subtitle="Hoàn thiện luồng quy trình phê duyệt, cấu hình ràng buộc kiểm tra biểu mẫu và mở rộng tập dữ liệu đa tầng cho Module Import Excel"
    )
    
    # --------------------------------------------------------------------------
    # I. THÔNG TIN CHUNG
    # --------------------------------------------------------------------------
    add_h1(doc, "I. THÔNG TIN CHUNG")
    add_header_info_table(doc)
    
    add_p(
        doc,
        "Tuần 4 là chặng tiếp nối quan trọng trong kỳ thực tập tốt nghiệp của em tại Khối Chính phủ số (eGov) thuộc Công ty TNHH Nền tảng chuyển đổi số Việt Nam. Trong tuần này, em tiếp tục thực hiện song song hai tuyến nhiệm vụ trọng tâm: (1) Tham gia giải quyết các lỗi nghiệp vụ và nâng cấp tiện ích giao diện cho dự án chính của công ty – Hệ thống báo cáo điện tử eForm; và (2) Rà soát, chuẩn bị mở rộng tập dữ liệu đa tầng phục vụ việc nghiên cứu phát triển đề tài cá nhân – Module Import file Excel thông minh, tự động ánh xạ dữ liệu.",
        bold_prefix="Bối cảnh thực tập:"
    )
    add_p(
        doc,
        "Căn cứ Kế hoạch học tập và Nhật ký thực tập hàng ngày, các nhiệm vụ cụ thể của em trong Tuần 4 bao gồm: Thứ 2 xử lý sự cố chuyển sai người nhận trong luồng trả lại báo cáo phê duyệt; Thứ 3 cấu hình logic kiểm tra và cảnh báo số liệu biểu mẫu tại tài khoản quản trị hệ thống; Thứ 4 kiểm tra và khắc phục lỗi xác thực (validate) khi nạp file Excel ở các biểu nhập tròn năm; Thứ 5 phát triển tính năng hiển thị toàn màn hình (Fullscreen) cho nội dung báo cáo; Thứ 6 rà soát dữ liệu đã chuẩn bị, bổ sung các trường hợp bảng nhiều tầng, nhiều bảng và dòng ghi chú phức tạp cho module Import Excel.",
        bold_prefix="Mục tiêu trọng tâm tuần 4:"
    )
    
    # --------------------------------------------------------------------------
    # II. NỘI DUNG VÀ KẾT QUẢ THỰC HIỆN CHI TIẾT
    # --------------------------------------------------------------------------
    add_h1(doc, "II. NỘI DUNG VÀ KẾT QUẢ THỰC HIỆN CHI TIẾT")
    
    # --------------------------------------------------------------------------
    # MỤC 1: QUÁ TRÌNH THỰC TẬP TẠI CÔNG TY
    # --------------------------------------------------------------------------
    add_h2(doc, "MỤC 1. QUÁ TRÌNH THỰC TẬP TẠI CÔNG TY: THỰC HIỆN CÁC NHIỆM VỤ DỰ ÁN CHÍNH (HỆ THỐNG EFORM)")
    
    # 1.1. Sửa chuyển sai người
    add_h3(doc, "1.1. Khảo sát luồng phê duyệt và sửa lỗi chuyển trả lại báo cáo sai người nhận (Thứ 2, 24/08/2026)")
    add_p(
        doc,
        "Trong quá trình kiểm thử luồng nghiệp vụ phê duyệt văn bản báo cáo trên hệ thống thực tế giữa các cấp đơn vị, nhóm kiểm thử và người dùng phát hiện một lỗi nghiêm trọng: Tại một đơn vị cụ thể do cán bộ Đặng Thị Thu Trang lập biểu và gửi báo cáo lên cấp trên là Cục Kế hoạch - Tài chính để kiểm tra (tài khoản cán bộ kiểm tra Cục Kế hoạch - Tài chính). Sau khi tiếp nhận báo cáo và nhận thấy số liệu cần đính chính, cán bộ kiểm duyệt thực hiện thao tác chuyển trả lại (trả lại báo cáo) để đơn vị cơ sở cập nhật lại số liệu. Tuy nhiên, trên hộp thoại xác nhận chuyển trả lại, hệ thống lại hiển thị và điều hướng trả về nhầm một tài khoản người nhận hoàn toàn xa lạ, dẫn tới việc cán bộ không thể chuyển trả lại báo cáo về đúng người khởi tạo ban đầu.",
        bold_prefix="Bối cảnh và sự cố nghiệp vụ phát sinh:"
    )
    add_p(
        doc,
        "Được anh Đào Đức Văn giao xử lý dứt điểm sự cố này, em đã tiến hành rà soát kỹ thuật theo quy trình chuyên sâu:",
        bold_prefix="Quy trình phân tích và khoanh vùng nguyên nhân:"
    )
    add_bullet(
        doc,
        "Em tiến hành kiểm tra mã nguồn tầng điều khiển và dịch vụ xử lý luồng chuyển tiếp văn bản Backend (.NET). Các điều kiện kiểm tra quyền và xác thực tham số đầu vào đều chuẩn chỉ, chứng tỏ mã nguồn xử lý không bị lỗi cú pháp hay ngoại lệ chương trình.",
        "Rà soát tầng xử lý Backend (.NET)"
    )
    add_bullet(
        doc,
        "Tiếp tục truy vết sâu vào cơ sở dữ liệu hệ thống tại các bảng định nghĩa quy trình (Workflow) và luồng trạng thái văn bản, em phát hiện nguyên nhân gốc rễ: Trong quá trình cấu hình hoặc cập nhật hệ thống trước đó, một liên kết nút (Transition Node) tại bước chuyển trả lại trong định nghĩa quy trình đã vô tình bị chỉnh sửa sai lệch mã định danh tài khoản đích (Target Node), làm đứt gãy quan hệ liên kết ngược về người lập biểu ban đầu.",
        "Phát hiện lỗi cấu hình nút quy trình (Workflow) trong CSDL"
    )
    add_p(
        doc,
        "Em đã tiến hành hiệu chỉnh lại cấu hình luồng quy trình (workflow) trực tiếp trong cơ sở dữ liệu hệ thống: thiết lập lại chính xác liên kết nút trả lại, đảm bảo đích chuyển tiếp luôn hướng về đúng tài khoản của cán bộ khởi tạo báo cáo ban đầu (chị Đặng Thị Thu Trang). Sau khi cập nhật, em đã kiểm thử lại toàn bộ chu trình nộp – duyệt – trả lại giữa tài khoản cơ sở và cán bộ kiểm tra Cục Kế hoạch - Tài chính; hệ thống đã hiển thị chuẩn xác người nhận và chuyển trạng thái báo cáo về đúng hộp thư xử lý của cán bộ lập biểu.",
        bold_prefix="Giải pháp xử lý và kết quả thực tế:"
    )
    add_image_figure(
        doc,
        "Docs/Tuần 4/01_cau_hinh_bieu_mau_admin.png",
        "Hình 1. Giao diện cấu hình quy trình xử lý (Workflow) và thiết lập các liên kết chuyển tiếp – trả lại báo cáo trên hệ thống eForm."
    )
    
    # 1.2. Cấu hình check logic
    add_h3(doc, "1.2. Cấu hình kiểm tra logic và thiết lập cảnh báo nghiệp vụ cho biểu mẫu ở tài khoản quản trị (Thứ 3, 25/08/2026)")
    add_p(
        doc,
        "Đối với hệ thống thông tin báo cáo thống kê phục vụ cơ quan nhà nước, tính nhất quán và chuẩn xác của số liệu là yếu tố sống còn. Để tránh tình trạng người dùng nhập số liệu tùy tiện, sai lệch giữa các chỉ tiêu tổng hợp và chỉ tiêu thành phần, hệ thống eForm yêu cầu phải thiết lập các quy tắc ràng buộc logic (Validation & Consistency Rules) ngay từ khâu cấu hình biểu mẫu trong phân hệ quản trị.",
        bold_prefix="Yêu cầu nghiệp vụ và vai trò của cấu hình logic:"
    )
    add_p(
        doc,
        "Em đã thực hiện nhiệm vụ cấu hình thông qua các bước tác nghiệp chuẩn mực trong phân hệ quản trị hệ thống:",
        bold_prefix="Quy trình thực hiện tại phân hệ quản trị:"
    )
    add_bullet(
        doc,
        "Đăng nhập vào hệ thống bằng tài khoản quản trị viên cấp cao, truy cập vào menu 'Danh mục' và chọn phân hệ 'Loại báo cáo số liệu' để quản lý toàn bộ danh sách các mẫu biểu thống kê.",
        "Truy cập danh mục quản trị biểu mẫu"
    )
    add_bullet(
        doc,
        "Sử dụng công cụ tìm kiếm lọc ra các biểu mẫu nghiệp vụ cần thiết lập; chọn biểu mẫu và mở tính năng 'Thiết kế biểu mẫu', sau đó chuyển sang thẻ 'Thiết kế kiểu dữ liệu'.",
        "Mở phân hệ Thiết kế kiểu dữ liệu"
    )
    add_bullet(
        doc,
        "Tại bảng cấu hình kiểu dữ liệu, em tiến hành rà soát và hiệu chỉnh cột 'Điều kiện': trực tiếp viết các biểu thức ràng buộc toán học và logic chuyên ngành (ví dụ: ràng buộc dòng tổng số bằng tổng các dòng con, giá trị cột chi tiết không được vượt quá cột tổng, hoặc tỷ lệ phần trăm phải tương ứng với tỷ số giữa tử số và mẫu số).",
        "Thiết lập công thức và điều kiện kiểm tra"
    )
    add_bullet(
        doc,
        "Thiết lập phân cấp cảnh báo trực quan: phân định rõ giữa các điều kiện bắt buộc (Error - khóa không cho gửi báo cáo khi vi phạm) và các điều kiện khuyến cáo (Warning - hiển thị thông báo nhắc nhở kiểm tra lại khi số liệu có dấu hiệu bất thường).",
        "Phân cấp thông điệp cảnh báo nghiệp vụ"
    )
    add_p(
        doc,
        "Sau khi hoàn tất cấu hình, các quy tắc kiểm tra logic đã phát huy hiệu quả ngay trên giao diện nhập biểu của cán bộ nghiệp vụ: khi nhập sai quy tắc, hệ thống tự động đổi màu cảnh báo tại ô vi phạm và đưa ra thông điệp hướng dẫn rõ ràng, giúp nâng cao đáng kể chất lượng dữ liệu đầu vào.",
        bold_prefix="Kết quả ghi nhận:"
    )
    add_image_figure(
        doc,
        "Docs/Tuần 4/01_danh_sach_loai_bao_cao_admin.png",
        "Hình 2. Danh mục quản lý Loại báo cáo số liệu trong phân hệ quản trị hệ thống eForm."
    )
    add_image_figure(
        doc,
        "Docs/Tuần 4/02_thiet_ke_kieu_du_lieu_va_dieu_kien.png",
        "Hình 3. Màn hình Thiết kế biểu mẫu và phân hệ Thiết kế kiểu dữ liệu hỗ trợ cấu hình cột Điều kiện kiểm tra logic."
    )
    
    # 1.3. Validate import tròn năm
    add_h3(doc, "1.3. Kiểm tra và khắc phục lỗi cơ chế xác thực (validate) khi nạp file Excel ở các biểu nhập tròn năm (Thứ 4, 26/08/2026)")
    add_p(
        doc,
        "Tính năng nạp dữ liệu từ file Excel (Import Excel) là công cụ được các cơ quan địa phương sử dụng thường xuyên để chuyển số liệu từ bảng tính có sẵn vào hệ thống eForm. Tuy nhiên, trong đợt kiểm thử với các mẫu biểu thuộc kỳ báo cáo tròn năm, nhóm nghiệp vụ phát hiện cơ chế kiểm tra tính hợp lệ dữ liệu (validate) hoàn toàn không hoạt động: tệp Excel dù có nhiều ô vi phạm công thức logic nhưng hệ thống vẫn cho phép nạp thành công mà không hiển thị bất kỳ cảnh báo nào.",
        bold_prefix="Hiện trạng lỗi phát sinh khi nạp file Excel:"
    )
    add_p(
        doc,
        "Em đã tiến hành kiểm tra sâu mã nguồn xử lý nạp tệp tại tầng giao diện (Frontend) và các tệp cấu hình quy tắc và phát hiện:",
        bold_prefix="Phân tích nguyên nhân kỹ thuật cốt lõi:"
    )
    add_bullet(
        doc,
        "Thuật toán phân tích điều kiện kiểm tra trước đây sử dụng cách tiếp cận dựa trên vị trí chỉ số mảng cố định: ngầm định cột điều kiện luôn nằm ở vị trí phần tử cuối cùng của mảng dữ liệu cấu hình.",
        "Thiết kế phụ thuộc chỉ số mảng tĩnh"
    )
    add_bullet(
        doc,
        "Gần đây, khi hệ thống bổ sung thêm tính năng cho phép chỉnh sửa trực tiếp trên ô dữ liệu (thuộc tính 'Allow edit'), cấu trúc mảng bị mở rộng và vị trí chỉ số của cột điều kiện bị xô lệch sang vị trí khác. Hậu quả là hàm đọc điều kiện đã trích xuất nhầm giá trị rỗng, dẫn tới việc bỏ qua toàn bộ vòng lặp kiểm tra quy tắc khi nạp bảng tính.",
        "Lệch vị trí chỉ số do bổ sung thuộc tính mới"
    )
    add_p(
        doc,
        "Em đã tái cấu trúc lại thuật toán đọc và kiểm tra điều kiện dữ liệu: chuyển đổi từ cơ chế trích xuất theo chỉ số mảng tương đối sang cơ chế định vị dựa trên tên khóa thuộc tính chuẩn hóa; đồng thời nâng cấp bộ phân tích công thức để nhận diện linh hoạt các biểu thức logic toán học (bao gồm các dạng so sánh nhỏ hơn hoặc bằng, chuỗi phép cộng các chỉ tiêu và công thức phần trăm). Em cũng bổ sung cơ chế hiển thị phản hồi trực quan: đánh dấu viền màu nổi bật tại các ô vi phạm điều kiện kèm bảng tổng hợp danh sách lỗi chi tiết để người dùng dễ dàng chỉnh sửa trước khi lưu chính thức.",
        bold_prefix="Giải pháp kỹ thuật đã hiện thực:"
    )
    add_p(
        doc,
        "Sau khi cập nhật mã nguồn, tính năng kiểm tra tính hợp lệ khi nạp dữ liệu Excel ở các biểu mẫu tròn năm đã hoạt động chuẩn xác 100%: hệ thống lập tức phát hiện các ô số liệu không khớp logic, hiển thị viền cảnh báo rõ ràng và ngăn chặn việc nạp dữ liệu sai vào hệ thống.",
        bold_prefix="Kết quả thực nghiệm:"
    )
    
    # 1.4. Fullscreen
    add_h3(doc, "1.4. Thiết kế và hiện thực chức năng toàn màn hình (Fullscreen) cho nội dung báo cáo (Thứ 5, 27/08/2026)")
    add_p(
        doc,
        "Hầu hết các mẫu biểu báo cáo chuyên ngành tư pháp có độ rộng rất lớn, thường từ 15 đến hơn 30 cột chỉ tiêu với nhiều tầng tiêu đề phức tạp. Trên bố cục giao diện mặc định, không gian làm việc của bảng tính bị giới hạn bởi thanh điều hướng bên trái và các khung thông tin bổ trợ, khiến cán bộ nhập liệu phải liên tục dùng chuột kéo thanh cuộn ngang và dọc, gây mỏi mắt và dễ nhầm lẫn dòng dữ liệu. Yêu cầu đặt ra là phải mở rộng tối đa không gian làm việc cho người dùng khi xem và nhập liệu báo cáo.",
        bold_prefix="Nhu cầu cấp thiết về trải nghiệm người dùng (UX):"
    )
    add_p(
        doc,
        "Em đã thiết kế và hiện thực giải pháp mở rộng không gian làm việc toàn màn hình (Fullscreen Overlay Architecture) với các đặc điểm kiến trúc tối ưu:",
        bold_prefix="Giải pháp thiết kế và kỹ thuật triển khai:"
    )
    add_bullet(
        doc,
        "Xây dựng một lớp phủ toàn màn hình (Overlay) với chỉ số phân lớp hiển thị cao, chiếm trọn 100% kích thước khung nhìn trình duyệt và tự động ẩn các thanh điều hướng tĩnh bên ngoài.",
        "Lớp phủ toàn màn hình chuyên dụng"
    )
    add_bullet(
        doc,
        "Tái cấu trúc bố cục khi chuyển sang chế độ toàn màn hình: toàn bộ thanh công cụ thao tác (Toolbar) chứa các nút lưu, gửi duyệt, xuất file được chuyển gọn gàng lên phần đầu trang (Header); vùng trung tâm dành trọn vẹn diện tích tối đa cho bảng tính lưới Handsontable; vùng ghi chú được bố trí linh hoạt ở chân trang (Footer).",
        "Tối ưu phân bổ không gian chức năng"
    )
    add_bullet(
        doc,
        "Tích hợp nút đóng trực quan tại góc trên cùng bên phải và hỗ trợ phím tắt tiêu chuẩn (Esc), cho phép người dùng chuyển đổi qua lại tức thì giữa chế độ toàn màn hình và chế độ xem thông thường mà không gây tải lại trang hay ảnh hưởng đến số liệu đang nhập dở dang.",
        "Cơ chế chuyển đổi trạng thái mượt mà"
    )
    add_p(
        doc,
        "Tính năng toàn màn hình sau khi hoàn thiện đã mang lại sự tiện lợi vượt bậc cho người sử dụng: không gian hiển thị bảng tính tăng hơn 65%, giúp cán bộ có thể bao quát toàn bộ các cột chỉ tiêu của mẫu biểu mà không bị gò bó tầm nhìn.",
        bold_prefix="Hiệu quả thực tế ghi nhận:"
    )
    add_image_figure(
        doc,
        "Docs/Tuần 4/03_giao_dien_nhap_lieu_va_fullscreen.png",
        "Hình 4. Giao diện mở rộng không gian làm việc toàn màn hình (Fullscreen) giúp tối ưu hóa diện tích hiển thị và nhập liệu bảng tính."
    )
    
    # --------------------------------------------------------------------------
    # MỤC 2: DỰ ÁN MODULE THỰC TẬP CÁ NHÂN
    # --------------------------------------------------------------------------
    add_h2(doc, "MỤC 2. DỰ ÁN THỰC TẬP CÁ NHÂN: NGHIÊN CỨU & XÂY DỰNG MODULE IMPORT EXCEL THÔNG MINH")
    
    # 2.1. Rà soát dữ liệu
    add_h3(doc, "2.1. Rà soát chuyên sâu bộ dữ liệu và bổ sung các trường hợp bảng biểu phức tạp (Thứ 6, 28/08/2026)")
    add_p(
        doc,
        "Tiếp nối kết quả phân loại họ bố cục (layout family) và xây dựng bộ quy tắc định chuẩn baseline ở Tuần 3, mục tiêu của Tuần 4 là nâng cao chất lượng và tính đại diện của bộ dữ liệu thực tế. Trong môi trường công tác thực tế tại các cơ quan nhà nước, các tệp bảng tính Excel thường không ở dạng chuẩn hóa lý tưởng mà chứa rất nhiều đặc thù dị biệt do thói quen soạn thảo văn bản hành chính.",
        bold_prefix="Mục tiêu mở rộng bộ dữ liệu thực tế:"
    )
    add_p(
        doc,
        "Em đã tiến hành rà soát chuyên sâu toàn bộ kho dữ liệu và thu thập bổ sung các trường hợp thử thách (Challenge Set):",
        bold_prefix="Các nhóm cấu trúc dị biệt được bổ sung:"
    )
    add_bullet(
        doc,
        "Các biểu mẫu có từ 4 đến 6 tầng tiêu đề lồng ghép phức tạp (như Mẫu biểu 04c, Biểu 21a). Tiêu đề cấp cha ôm trọn nhiều nhóm tiêu đề con bên dưới thông qua các khối ô gộp ngang và dọc, đòi hỏi thuật toán phải bóc tách chính xác mối quan hệ phân cấp để không làm xô lệch các cột số liệu.",
        "Bảng biểu có cấu trúc tiêu đề đa tầng"
    )
    add_bullet(
        doc,
        "Các tệp Excel mà trên một trang tính (sheet) bố trí từ 2 đến 3 bảng số liệu độc lập nằm kế tiếp nhau theo chiều dọc hoặc song song (như Mẫu biểu 01d). Mô hình cần nhận diện được ranh giới ngăn cách giữa các bảng để không gộp nhầm dữ liệu.",
        "Trang tính chứa nhiều vùng bảng độc lập"
    )
    add_bullet(
        doc,
        "Nhiều biểu mẫu có các dòng ghi chú, căn cứ ban hành, dòng chỉ dẫn công thức hoặc khu vực chữ ký xác nhận của thủ trưởng cơ quan nằm sát ngay bên dưới dòng tổng số. Đây là nguyên nhân hàng đầu khiến các bộ parser truyền thống bóc tách nhầm chữ ký thành dữ liệu nghiệp vụ.",
        "Bảng biểu có chú thích và khu vực chữ ký sát mép bảng"
    )
    add_image_figure(
        doc,
        "Docs/Tuần 4/diagrams/02_QUY_TRINH_XAY_DUNG_BO_DU_LIEU.png",
        "Hình 5. Sơ đồ quy trình chuẩn bị, làm sạch, gắn nhãn và phân bổ bộ dữ liệu thử nghiệm đa tầng cho module Import Excel."
    )
    
    # 2.2. Biểu diễn cấu trúc tiêu đề
    add_h3(doc, "2.2. Biểu diễn cấu trúc tiêu đề phân cấp và hoàn thiện dữ liệu kiểm thử")
    add_p(
        doc,
        "Để chuyển đổi dữ liệu cấu trúc bảng tính sang định dạng chuẩn hóa sẵn sàng cho việc huấn luyện mô hình học máy và trực quan hóa lên giao diện người dùng, em đã thiết kế giải pháp biểu diễn đường dẫn tiêu đề phân cấp (Header Path Representation):",
        bold_prefix="Phương pháp biểu diễn cấu trúc:"
    )
    add_bullet(
        doc,
        "Dựa vào ma trận hình học các ô gộp, thuật toán tự động duyệt từ tầng tiêu đề cao nhất xuống tầng tiêu đề chi tiết để tạo ra một chuỗi phân cấp hoàn chỉnh cho từng cột số liệu (ví dụ: 'Văn bản QPPL > Thẩm định > Số lượng').",
        "Thiết lập đường dẫn phân cấp chỉ tiêu"
    )
    add_bullet(
        doc,
        "Tách riêng dòng mã ký hiệu cột (như (1), (2), A, B...) khỏi phần nội dung nhãn chỉ tiêu, tạo điều kiện thuận lợi cho bước tự động ánh xạ sang các trường dữ liệu đích của hệ thống eForm sau này.",
        "Chuẩn hóa mã cột và nhãn số liệu"
    )
    add_bullet(
        doc,
        "Bộ dữ liệu sau khi mở rộng đã được kiểm tra tính toàn vẹn kỹ thuật: 100% mẫu có mã định danh duy nhất, ranh giới vùng bảng nằm hoàn toàn trong kích thước trang tính, các vùng bảng không chồng lấn và toàn bộ dữ liệu thuộc cùng một họ bố cục được bảo toàn trong duy nhất một phân tập đánh giá.",
        "Kiểm định tính toàn vẹn kỹ thuật của bộ dữ liệu"
    )
    add_image_figure(
        doc,
        "Docs/Tuần 4/diagrams/03_BIEU_DIEN_TIEU_DE_PHAN_CAP.png",
        "Hình 6. Biểu diễn cấu trúc tiêu đề phân cấp đa tầng và ranh giới vùng bảng tính phục vụ huấn luyện mô hình."
    )
    
    # --------------------------------------------------------------------------
    # III. BẢNG TỔNG HỢP TIẾN ĐỘ VÀ KẾT QUẢ ĐẠT ĐƯỢC
    # --------------------------------------------------------------------------
    add_h1(doc, "III. BẢNG TỔNG HỢP TIẾN ĐỘ VÀ KẾT QUẢ ĐẠT ĐƯỢC THEO NHẬT KÝ THỰC TẬP")
    add_p(
        doc,
        "Dưới đây là bảng đối chiếu chi tiết giữa nội dung công việc được phân công và kết quả thực tế em đã hoàn thành trong Tuần 4, bám sát từng ngày theo đúng Nhật ký thực tập hàng ngày đã được cán bộ hướng dẫn phê duyệt:"
    )
    
    headers_w4 = ["STT", "Thời gian", "Nhiệm vụ công việc thực hiện", "Kết quả thực tế ghi nhận", "Đánh giá"]
    rows_w4 = [
        [
            "1",
            "Thứ 2\n24/08/2026",
            "Sửa lại chuyển báo cáo đang chuyển sai người nhận trong luồng phê duyệt.",
            "Rà soát backend và CSDL, sửa lại cấu hình node trong workflow; chức năng trả lại báo cáo đã chuyển đúng đối tượng.",
            "Hoàn thành"
        ],
        [
            "2",
            "Thứ 3\n25/08/2026",
            "Cấu hình check logic kiểm tra biểu mẫu ở tài khoản quản trị hệ thống.",
            "Đã thiết lập đầy đủ các quy tắc logic toán học và điều kiện cảnh báo số liệu tại cột Điều kiện trong phân hệ admin.",
            "Hoàn thành"
        ],
        [
            "3",
            "Thứ 4\n26/08/2026",
            "Check lại validate khi import dữ liệu ở các biểu nhập tròn năm.",
            "Khắc phục triệt để lỗi lệch chỉ số cột điều kiện; chức năng kiểm tra hợp lệ khi nạp file Excel hoạt động chuẩn xác.",
            "Hoàn thành"
        ],
        [
            "4",
            "Thứ 5\n27/08/2026",
            "Thêm chức năng Fullscreen cho nội dung biểu mẫu báo cáo.",
            "Hiện thực thành công lớp phủ toàn màn hình; giao diện mở rộng tối đa 100% diện tích làm việc, hỗ trợ phím tắt Esc.",
            "Hoàn thành"
        ],
        [
            "5",
            "Thứ 6\n28/08/2026",
            "Rà soát dữ liệu đã chuẩn bị, bổ sung các trường hợp bảng nhiều tầng, nhiều bảng và dòng ghi chú. (Import Excel)",
            "Bộ dữ liệu thử nghiệm bao quát hơn, xử lý tốt các tình huống biên phức tạp, sẵn sàng cho bước huấn luyện mô hình.",
            "Hoàn thành"
        ]
    ]
    add_custom_table(doc, headers_w4, rows_w4, [Inches(0.6), Inches(1.2), Inches(2.3), Inches(2.1), Inches(0.8)])
    
    # --------------------------------------------------------------------------
    # IV. KHÓ KHĂN GẶP PHẢI VÀ BÀI HỌC KINH NGHIỆM
    # --------------------------------------------------------------------------
    add_h1(doc, "IV. KHÓ KHĂN GẶP PHẢI VÀ BÀI HỌC KINH NGHIỆM")
    add_bullet(
        doc,
        "Hệ thống quản lý quy trình (Workflow Engine) của các nền tảng Chính phủ số có độ phức tạp rất cao với hàng chục bước chuyển tiếp có điều kiện. Việc can thiệp trực tiếp vào bảng cấu hình quy trình trong cơ sở dữ liệu đòi hỏi phải đọc hiểu tường tận mô hình quan hệ giữa các bảng trạng thái, bảng vết xử lý và bảng phân quyền người dùng để không làm phát sinh lỗi dây chuyền sang các quy trình khác.",
        "Kỹ năng kiểm soát luồng quy trình nghiệp vụ phức tạp"
    )
    add_bullet(
        doc,
        "Lỗi mất khả năng kiểm tra hợp lệ khi thêm thuộc tính 'Allow edit' là một bài học đắt giá về kiến trúc phần mềm: việc sử dụng các chỉ số mảng cố định (Hardcoded Array Index) rất dễ bị phá vỡ khi hệ thống nâng cấp. Em đã rút ra kinh nghiệm sâu sắc về việc luôn luôn truy xuất dữ liệu dựa trên khóa thuộc tính định danh để đảm bảo tính bền vững của mã nguồn.",
        "Tránh phụ thuộc vào vị trí mảng tĩnh trong mở rộng phần mềm"
    )
    add_bullet(
        doc,
        "Các biểu mẫu hành chính có phần ghi chú và chữ ký nằm sát bảng tính là thách thức lớn đối với bài toán nhận diện thị giác và cấu trúc. Việc xây dựng kịch bản kiểm thử độc lập cho từng cạnh ranh giới (trên, dưới, trái, phải) đã giúp em có cái nhìn toàn diện hơn về bài toán xử lý dữ liệu thực tế.",
        "Xử lý bài toán ranh giới dữ liệu nhiễu"
    )
    
    # --------------------------------------------------------------------------
    # V. KẾT LUẬN VÀ KẾ HOẠCH TUẦN TIẾP THEO
    # --------------------------------------------------------------------------
    add_h1(doc, "V. KẾT LUẬN VÀ KẾ HOẠCH TUẦN TIẾP THEO")
    add_p(
        doc,
        "Trong Tuần 4, em đã hoàn thành 100% mục tiêu công việc đề ra theo đúng kế hoạch học tập. Các vấn đề kỹ thuật trọng tâm của dự án chính eForm (sửa luồng trả lại báo cáo, cấu hình logic kiểm tra, sửa lỗi xác thực khi nạp Excel và nâng cấp toàn màn hình) đều được nghiệm thu thành công. Đồng thời, bộ dữ liệu thử nghiệm cho module cá nhân đã được mở rộng hoàn chỉnh, tạo tiền đề vững chắc cho giai đoạn huấn luyện mô hình.",
        bold_prefix="Đánh giá kết quả tuần 4:"
    )
    add_p(
        doc,
        "Bám sát tiến độ trong Nhật ký thực tập hàng ngày, các nhiệm vụ trọng tâm em sẽ thực hiện trong Tuần 5 bao gồm:",
        bold_prefix="Kế hoạch công tác Tuần 5 (Từ 31/08/2026 đến 04/09/2026):"
    )
    add_bullet(
        doc,
        "Thứ 2 (31/08/2026): Họp bàn về tiến độ dự án eForm và tiếp nhận các đầu việc kỹ thuật tuần 5.",
        "Dự án eForm"
    )
    add_bullet(
        doc,
        "Thứ 3 (01/09/2026): Khảo sát và tổng hợp lại dữ liệu bị lỗi #REF tại các dòng tính tổng, khắc phục sự cố mất dữ liệu của đơn vị cấp xã và cơ quan ngang bộ.",
        "Dự án eForm"
    )
    add_bullet(
        doc,
        "Thứ 4 (02/09/2026): Sửa lỗi đối với biểu 11b khi mở 2 báo cáo con cùng lúc dẫn tới báo cáo ban đầu không hiển thị phần ghi chú.",
        "Dự án eForm"
    )
    add_bullet(
        doc,
        "Thứ 5 (03/09/2026): Khắc phục lỗi khi nhấn tổng hợp lại dữ liệu thì bị mất dữ liệu nhập tay; xây dựng cơ chế tự động lưu dữ liệu an toàn.",
        "Dự án eForm"
    )
    add_bullet(
        doc,
        "Thứ 6 (04/09/2026): Huấn luyện thử nghiệm mô hình nhận diện cấu trúc bảng tính Excel và chuẩn hóa kết quả đầu ra theo định dạng thống nhất.",
        "Module Import Excel"
    )
    
    out_path = r"c:\Users\hoang\Downloads\eform_btp\Docs\Tuần 4\BÁO cáo thực tập tuần 4.docx"
    doc.save(out_path)
    print(f"Đã tạo thành công báo cáo thực tập: {out_path}")

if __name__ == "__main__":
    generate_week_4()
