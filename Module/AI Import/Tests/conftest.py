from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font


@pytest.fixture()
def sample_workbook(tmp_path: Path) -> Path:
    path = tmp_path / "bao-cao-thu-nghiem.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Báo cáo"
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "BÁO CÁO TỔNG HỢP"
    for cell, value in zip(sheet[3], ["STT", "Tên đơn vị", "Số lượng"]):
        cell.value = value
        cell.font = Font(bold=True)
    sheet.append([1, "Đơn vị A", 10])
    sheet.append([2, "Đơn vị B", 20])
    hidden = workbook.create_sheet("Danh mục")
    hidden.sheet_state = "hidden"
    hidden.append(["Mã", "Tên"])
    workbook.save(path)
    workbook.close()
    return path

