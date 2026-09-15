from __future__ import annotations

import argparse
import copy
import hashlib
import html
import json
import math
import os
import random
import re
import tempfile
import unicodedata
from collections import defaultdict
from datetime import date, datetime, time, timezone
from pathlib import Path
from typing import Any

os.environ.setdefault("MPLCONFIGDIR", str(Path(tempfile.gettempdir()) / "eform-matplotlib"))
os.environ["MPLBACKEND"] = "Agg"
import torch
import torch.nn as nn
import torch.nn.functional as F
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from torch.utils.data import DataLoader, Dataset
from tqdm.auto import tqdm


MODEL_NAME = "eform_excel_table_extractor_v1"
MODEL_FILENAME = f"{MODEL_NAME}.eformmodel"
FORMAT_VERSION = "eformmodel-1.0"
CLASS_NAMES = ("outside", "header", "value")
IGNORE_INDEX = -100
TEXT_HASH_DIM = 14
FEATURE_CHANNELS = 18 + TEXT_HASH_DIM
NUMBER_IN_TEXT_RE = re.compile(r"\d+(?:[.,]\d+)*")
FOOTER_TEXT_RE = re.compile(
    r"\b(ghi chu|nguoi lap bieu|nguoi kiem tra|ky,|dong dau|thu truong|giam doc)\b"
)


def seed_everything(seed: int) -> None:
    random.seed(seed)
    torch.manual_seed(seed)
    if torch.cuda.is_available():
        torch.cuda.manual_seed_all(seed)
    torch.backends.cudnn.deterministic = True
    torch.backends.cudnn.benchmark = False


def json_safe(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else str(value)
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    return str(value)


def clean_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value if value is not None else "")).strip()


def fold_text(value: Any) -> str:
    text = unicodedata.normalize("NFD", clean_text(value).lower())
    text = "".join(character for character in text if unicodedata.category(character) != "Mn")
    return re.sub(r"\s+", " ", text.replace("đ", "d")).strip()


def _cell_value(worksheet: Any, row: int, column: int) -> Any:
    cell = worksheet.cell(row, column)
    return None if isinstance(cell, MergedCell) else cell.value


def _integer(value: Any) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str) and re.fullmatch(r"[-+]?\d+", value.strip()):
        return int(value.strip())
    return None


def _value_kind(value: Any, data_type: str) -> str:
    if value is None or (isinstance(value, str) and not value.strip()):
        return "blank"
    if data_type == "f" or (isinstance(value, str) and value.startswith("=")):
        return "formula"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (date, datetime, time)):
        return "temporal"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return "numeric"
    return "text"


def _text_hash(text: str) -> list[float]:
    vector = [0.0] * TEXT_HASH_DIM
    normalized = NUMBER_IN_TEXT_RE.sub("<N>", fold_text(text))[:256]
    if not normalized:
        return vector
    padded = f"^{normalized}$"
    grams = [padded[index:index + 3] for index in range(max(1, len(padded) - 2))]
    for gram in grams:
        digest = hashlib.blake2b(gram.encode("utf-8"), digest_size=4).digest()
        number = int.from_bytes(digest, "little")
        vector[number % TEXT_HASH_DIM] += -1.0 if number & 1 else 1.0
    scale = math.sqrt(max(1, len(grams)))
    return [value / scale for value in vector]


def actual_bounds(worksheet: Any) -> tuple[int, int]:
    max_row = 0
    max_column = 0
    for cell in getattr(worksheet, "_cells", {}).values():
        value = getattr(cell, "value", None)
        if value is not None and (not isinstance(value, str) or value.strip()):
            max_row = max(max_row, int(cell.row))
            max_column = max(max_column, int(cell.column))
    for merged_range in worksheet.merged_cells.ranges:
        left, top, right, bottom = range_boundaries(str(merged_range))
        max_row = max(max_row, bottom)
        max_column = max(max_column, right)
    return max_row, max_column


def encode_sheet(worksheet: Any, rows: int, columns: int) -> torch.Tensor:
    """Turn an Excel sheet into a value-safe CxHxW tensor.

    Numeric/date/formula contents are represented only by type. Text is mapped
    through deterministic hashing, so the network cannot reproduce cell values.
    """
    features = torch.zeros((FEATURE_CHANNELS, rows, columns), dtype=torch.float32)
    if rows == 0 or columns == 0:
        return features

    for row in range(1, rows + 1):
        numeric_values = [
            number
            for column in range(1, columns + 1)
            if (number := _integer(_cell_value(worksheet, row, column))) is not None
        ]
        negative_sequence = len(numeric_values) >= 2 and numeric_values == list(
            range(-1, -len(numeric_values) - 1, -1)
        )
        for column in range(1, columns + 1):
            cell = worksheet.cell(row, column)
            value = _cell_value(worksheet, row, column)
            kind = _value_kind(value, getattr(cell, "data_type", "") or "")
            row_index, column_index = row - 1, column - 1
            features[0, row_index, column_index] = float(kind != "blank")
            features[1, row_index, column_index] = float(kind == "text")
            features[2, row_index, column_index] = float(kind == "numeric")
            features[3, row_index, column_index] = float(kind == "formula")
            features[4, row_index, column_index] = float(kind == "temporal")
            features[5, row_index, column_index] = float(kind == "boolean")
            features[6, row_index, column_index] = float(bool(getattr(cell.font, "bold", False)))
            features[7, row_index, column_index] = float(
                getattr(cell.alignment, "horizontal", None) in {"center", "centerContinuous"}
            )
            features[8, row_index, column_index] = sum(
                bool(getattr(getattr(cell.border, side, None), "style", None))
                for side in ("left", "right", "top", "bottom")
            ) / 4.0
            features[13, row_index, column_index] = row_index / max(1, rows - 1)
            features[14, row_index, column_index] = column_index / max(1, columns - 1)
            integer = _integer(value)
            features[15, row_index, column_index] = float(integer is not None and integer < 0)
            features[16, row_index, column_index] = float(negative_sequence)
            text = clean_text(value) if kind == "text" else ""
            features[17, row_index, column_index] = min(len(text), 200) / 200.0
            if text:
                features[18:, row_index, column_index] = torch.tensor(_text_hash(text))

    for merged_range in worksheet.merged_cells.ranges:
        left, top, right, bottom = range_boundaries(str(merged_range))
        left, top = max(1, left), max(1, top)
        right, bottom = min(columns, right), min(rows, bottom)
        if left > right or top > bottom:
            continue
        features[9, top - 1:bottom, left - 1:right] = 1.0
        features[10, top - 1, left - 1] = 1.0
        features[11, top - 1:bottom, left - 1:right] = min(bottom - top + 1, 20) / 20.0
        features[12, top - 1:bottom, left - 1:right] = min(right - left + 1, 20) / 20.0
    return features


def target_from_tables(rows: int, columns: int, tables: list[dict[str, int]]) -> torch.Tensor:
    target = torch.zeros((rows, columns), dtype=torch.long)
    for table in tables:
        top, left = table["top"] - 1, table["left"] - 1
        bottom, right = table["bottom"], table["right"]
        header_bottom = table["header_bottom"]
        target[top:header_bottom, left:right] = 1
        if header_bottom < bottom:
            target[header_bottom:bottom, left:right] = 2
    return target


def apply_augmentation(
    features: torch.Tensor,
    target: torch.Tensor,
    recipe: dict[str, Any] | None,
) -> tuple[torch.Tensor, torch.Tensor]:
    if not recipe:
        return features, target
    row_shift = int(recipe.get("row_shift", 0))
    column_shift = int(recipe.get("column_shift", 0))
    channels, rows, columns = features.shape
    shifted_features = torch.zeros(
        (channels, rows + row_shift, columns + column_shift), dtype=features.dtype
    )
    shifted_target = torch.zeros((rows + row_shift, columns + column_shift), dtype=target.dtype)
    shifted_features[:, row_shift:, column_shift:] = features
    shifted_target[row_shift:, column_shift:] = target

    generator = torch.Generator().manual_seed(int(recipe.get("seed", 0)))
    style_dropout = float(recipe.get("style_dropout", 0.0))
    if style_dropout:
        style_mask = torch.rand(
            shifted_target.shape, generator=generator
        ) < style_dropout
        shifted_features[6:9, style_mask] = 0.0
    header_dropout = float(recipe.get("header_text_dropout", 0.0))
    if header_dropout:
        text_mask = (
            torch.rand(shifted_target.shape, generator=generator) < header_dropout
        ) & (shifted_target == 1)
        shifted_features[17:, text_mask] = 0.0
    return shifted_features, shifted_target


class ExcelGridDataset(Dataset):
    def __init__(
        self,
        module_root: Path,
        records: list[dict[str, Any]],
        include_augmentations: bool,
    ) -> None:
        self.module_root = module_root
        self.records = records
        self.items: list[tuple[int, dict[str, Any] | None]] = []
        self.cache: dict[str, tuple[torch.Tensor, torch.Tensor]] = {}
        for record_index, record in enumerate(records):
            self.items.append((record_index, None))
            if include_augmentations:
                self.items.extend(
                    (record_index, recipe)
                    for recipe in record.get("augmentation_recipes", [])
                )

    def __len__(self) -> int:
        return len(self.items)

    def _base(self, record: dict[str, Any]) -> tuple[torch.Tensor, torch.Tensor]:
        sample_id = record["sample_id"]
        if sample_id not in self.cache:
            workbook = load_workbook(
                self.module_root / record["source_file"],
                read_only=False,
                data_only=False,
                keep_links=False,
            )
            try:
                worksheet = workbook.worksheets[record["sheet_index"]]
                rows, columns = record["sheet_shape"]
                features = encode_sheet(worksheet, rows, columns)
                target = target_from_tables(rows, columns, record["tables"])
            finally:
                workbook.close()
            self.cache[sample_id] = features, target
        return self.cache[sample_id]

    def __getitem__(self, index: int) -> tuple[torch.Tensor, torch.Tensor, dict[str, Any]]:
        record_index, recipe = self.items[index]
        record = self.records[record_index]
        features, target = self._base(record)
        features, target = apply_augmentation(features, target, recipe)
        metadata = {
            "sample_id": record["sample_id"],
            "source_file": record["source_file"],
            "sheet_index": record["sheet_index"],
            "tables": record["tables"],
            "split": record["split"],
            "augmented": recipe is not None,
        }
        if recipe:
            row_shift = int(recipe.get("row_shift", 0))
            column_shift = int(recipe.get("column_shift", 0))
            metadata["tables"] = [
                {
                    "top": table["top"] + row_shift,
                    "left": table["left"] + column_shift,
                    "bottom": table["bottom"] + row_shift,
                    "right": table["right"] + column_shift,
                    "header_bottom": table["header_bottom"] + row_shift,
                }
                for table in record["tables"]
            ]
        return features, target, metadata


def collate_grids(batch: list[tuple[torch.Tensor, torch.Tensor, dict[str, Any]]]) -> dict[str, Any]:
    max_rows = max(features.shape[1] for features, _, _ in batch)
    max_columns = max(features.shape[2] for features, _, _ in batch)
    inputs = torch.zeros((len(batch), FEATURE_CHANNELS, max_rows, max_columns), dtype=torch.float32)
    targets = torch.full((len(batch), max_rows, max_columns), IGNORE_INDEX, dtype=torch.long)
    shapes: list[tuple[int, int]] = []
    metadata: list[dict[str, Any]] = []
    for index, (features, target, item_metadata) in enumerate(batch):
        rows, columns = target.shape
        inputs[index, :, :rows, :columns] = features
        targets[index, :rows, :columns] = target
        shapes.append((rows, columns))
        metadata.append(item_metadata)
    return {"inputs": inputs, "targets": targets, "shapes": shapes, "metadata": metadata}


def _group_count(channels: int) -> int:
    for groups in (8, 4, 2):
        if channels % groups == 0:
            return groups
    return 1


class ConvBlock(nn.Module):
    def __init__(self, input_channels: int, output_channels: int) -> None:
        super().__init__()
        self.block = nn.Sequential(
            nn.Conv2d(input_channels, output_channels, 3, padding=1, bias=False),
            nn.GroupNorm(_group_count(output_channels), output_channels),
            nn.GELU(),
            nn.Conv2d(output_channels, output_channels, 3, padding=1, bias=False),
            nn.GroupNorm(_group_count(output_channels), output_channels),
            nn.GELU(),
        )

    def forward(self, inputs: torch.Tensor) -> torch.Tensor:
        return self.block(inputs)


class EFormTableNet(nn.Module):
    """A compact U-Net that classifies each cell as outside/header/value."""

    def __init__(self, input_channels: int = FEATURE_CHANNELS, base_channels: int = 32) -> None:
        super().__init__()
        self.encoder1 = ConvBlock(input_channels, base_channels)
        self.encoder2 = ConvBlock(base_channels, base_channels * 2)
        self.bottleneck = ConvBlock(base_channels * 2, base_channels * 3)
        self.decoder2 = ConvBlock(base_channels * 5, base_channels * 2)
        self.decoder1 = ConvBlock(base_channels * 3, base_channels)
        self.head = nn.Conv2d(base_channels, len(CLASS_NAMES), 1)

    def forward(self, inputs: torch.Tensor) -> torch.Tensor:
        encoder1 = self.encoder1(inputs)
        encoder2 = self.encoder2(F.max_pool2d(encoder1, 2, ceil_mode=True))
        bottleneck = self.bottleneck(F.max_pool2d(encoder2, 2, ceil_mode=True))
        decoder2 = F.interpolate(bottleneck, size=encoder2.shape[-2:], mode="bilinear", align_corners=False)
        decoder2 = self.decoder2(torch.cat((decoder2, encoder2), dim=1))
        decoder1 = F.interpolate(decoder2, size=encoder1.shape[-2:], mode="bilinear", align_corners=False)
        decoder1 = self.decoder1(torch.cat((decoder1, encoder1), dim=1))
        return self.head(decoder1)


def update_confusion(confusion: torch.Tensor, predictions: torch.Tensor, targets: torch.Tensor) -> None:
    valid = targets != IGNORE_INDEX
    encoded = targets[valid] * len(CLASS_NAMES) + predictions[valid]
    confusion += torch.bincount(encoded.cpu(), minlength=len(CLASS_NAMES) ** 2).reshape(
        len(CLASS_NAMES), len(CLASS_NAMES)
    )


def metrics_from_confusion(confusion: torch.Tensor) -> dict[str, float]:
    matrix = confusion.to(torch.float64)
    total = matrix.sum().clamp_min(1)
    metrics: dict[str, float] = {"cell_accuracy": float(matrix.diag().sum() / total)}
    f1_scores: list[float] = []
    for index, name in enumerate(CLASS_NAMES):
        true_positive = matrix[index, index]
        false_positive = matrix[:, index].sum() - true_positive
        false_negative = matrix[index, :].sum() - true_positive
        precision = true_positive / (true_positive + false_positive).clamp_min(1)
        recall = true_positive / (true_positive + false_negative).clamp_min(1)
        f1 = 2 * precision * recall / (precision + recall).clamp_min(1e-12)
        union = true_positive + false_positive + false_negative
        iou = true_positive / union.clamp_min(1)
        metrics[f"{name}_precision"] = float(precision)
        metrics[f"{name}_recall"] = float(recall)
        metrics[f"{name}_f1"] = float(f1)
        metrics[f"{name}_iou"] = float(iou)
        f1_scores.append(float(f1))
    foreground_true_positive = matrix[1:, 1:].sum()
    foreground_target = matrix[1:, :].sum()
    foreground_prediction = matrix[:, 1:].sum()
    foreground_union = foreground_target + foreground_prediction - foreground_true_positive
    metrics["foreground_iou"] = float(foreground_true_positive / foreground_union.clamp_min(1))
    metrics["macro_f1"] = sum(f1_scores) / len(f1_scores)
    metrics["mean_header_value_iou"] = (metrics["header_iou"] + metrics["value_iou"]) / 2
    return metrics


def prediction_to_tables(
    prediction: torch.Tensor,
    probabilities: torch.Tensor | None = None,
    minimum_component_cells: int = 4,
    header_fraction_threshold: float = 0.30,
) -> list[dict[str, Any]]:
    prediction = prediction.cpu()
    rows, columns = prediction.shape
    foreground = prediction != 0
    visited = torch.zeros_like(foreground, dtype=torch.bool)
    components: list[list[tuple[int, int]]] = []
    for row in range(rows):
        for column in range(columns):
            if not foreground[row, column] or visited[row, column]:
                continue
            stack = [(row, column)]
            visited[row, column] = True
            component: list[tuple[int, int]] = []
            while stack:
                current_row, current_column = stack.pop()
                component.append((current_row, current_column))
                for next_row, next_column in (
                    (current_row - 1, current_column),
                    (current_row + 1, current_column),
                    (current_row, current_column - 1),
                    (current_row, current_column + 1),
                ):
                    if (
                        0 <= next_row < rows
                        and 0 <= next_column < columns
                        and foreground[next_row, next_column]
                        and not visited[next_row, next_column]
                    ):
                        visited[next_row, next_column] = True
                        stack.append((next_row, next_column))
            if len(component) >= minimum_component_cells:
                components.append(component)

    tables: list[dict[str, Any]] = []
    for component in components:
        component_rows = [row for row, _ in component]
        component_columns = [column for _, column in component]
        top, bottom = min(component_rows), max(component_rows)
        left, right = min(component_columns), max(component_columns)
        header_rows = []
        for row in range(top, bottom + 1):
            header_fraction = float((prediction[row, left:right + 1] == 1).float().mean())
            if header_fraction >= header_fraction_threshold:
                header_rows.append(row)
        header_bottom = max(header_rows, default=top)
        confidence = None
        if probabilities is not None:
            confidence = float(
                torch.stack([probabilities[prediction[row, column], row, column] for row, column in component]).mean()
            )
        tables.append({
            "top": top + 1,
            "left": left + 1,
            "bottom": bottom + 1,
            "right": right + 1,
            "header_bottom": header_bottom + 1,
            "confidence": confidence,
        })
    tables.sort(key=lambda table: (table["top"], table["left"], table["bottom"], table["right"]))
    return tables


def refine_table_bounds(worksheet: Any, table: dict[str, Any]) -> dict[str, Any]:
    """Trim obvious note/signature footers accidentally connected to a predicted table."""
    refined = dict(table)
    left, right = int(refined["left"]), int(refined["right"])
    header_bottom = int(refined["header_bottom"])
    width = max(1, right - left + 1)
    bottom = int(refined["bottom"])
    trimmed = 0

    while bottom > header_bottom:
        cells = [worksheet.cell(bottom, column) for column in range(left, right + 1)]
        values = [_cell_value(worksheet, bottom, column) for column in range(left, right + 1)]
        nonempty = sum(bool(clean_text(value)) for value in values)
        numeric = sum(
            isinstance(value, (int, float)) and not isinstance(value, bool)
            for value in values
        )
        bordered = sum(
            any(getattr(cell.border, side).style for side in ("left", "right", "top", "bottom"))
            for cell in cells
        )
        text = fold_text(" ".join(clean_text(value) for value in values if value is not None))
        sparse = nonempty <= max(1, math.ceil(width * 0.30))
        weak_grid = bordered < math.ceil(width * 0.35)
        footer_text = bool(FOOTER_TEXT_RE.search(text))
        if numeric == 0 and weak_grid and (sparse or footer_text):
            bottom -= 1
            trimmed += 1
            continue
        break

    if bottom > header_bottom:
        refined["bottom"] = bottom
        if trimmed:
            refined["postprocessing"] = "footer-trim-v1"
            refined["trimmed_footer_rows"] = trimmed
    return refined


def box_iou(first: dict[str, Any], second: dict[str, Any]) -> float:
    intersection_top = max(first["top"], second["top"])
    intersection_left = max(first["left"], second["left"])
    intersection_bottom = min(first["bottom"], second["bottom"])
    intersection_right = min(first["right"], second["right"])
    intersection = max(0, intersection_bottom - intersection_top + 1) * max(
        0, intersection_right - intersection_left + 1
    )
    first_area = (first["bottom"] - first["top"] + 1) * (first["right"] - first["left"] + 1)
    second_area = (second["bottom"] - second["top"] + 1) * (second["right"] - second["left"] + 1)
    return intersection / max(1, first_area + second_area - intersection)


def compare_table_sets(predicted: list[dict[str, Any]], expected: list[dict[str, Any]]) -> dict[str, float]:
    if not predicted and not expected:
        return {"box_iou": 1.0, "box_iou_at_50": 1.0, "exact_region": 1.0}
    available = set(range(len(predicted)))
    matched_ious: list[float] = []
    exact_count = 0
    for gold in expected:
        if not available:
            matched_ious.append(0.0)
            continue
        best_index = max(available, key=lambda index: box_iou(predicted[index], gold))
        score = box_iou(predicted[best_index], gold)
        matched_ious.append(score)
        candidate = predicted[best_index]
        exact_count += int(all(candidate[key] == gold[key] for key in ("top", "left", "bottom", "right", "header_bottom")))
        available.remove(best_index)
    denominator = max(1, len(expected), len(predicted))
    return {
        "box_iou": sum(matched_ious) / denominator,
        "box_iou_at_50": sum(score >= 0.50 for score in matched_ious) / denominator,
        "exact_region": float(len(predicted) == len(expected) and exact_count == len(expected)),
    }


def run_epoch(
    model: nn.Module,
    loader: DataLoader,
    criterion: nn.Module,
    device: torch.device,
    optimizer: torch.optim.Optimizer | None,
    description: str,
    gradient_clip_norm: float = 2.0,
    minimum_component_cells: int = 4,
    header_fraction_threshold: float = 0.30,
) -> dict[str, float]:
    training = optimizer is not None
    model.train(training)
    confusion = torch.zeros((len(CLASS_NAMES), len(CLASS_NAMES)), dtype=torch.long)
    total_loss = 0.0
    box_totals: defaultdict[str, float] = defaultdict(float)
    sample_count = 0
    progress = tqdm(loader, desc=description, leave=False)
    for batch_index, batch in enumerate(progress, 1):
        inputs = batch["inputs"].to(device, non_blocking=True)
        targets = batch["targets"].to(device, non_blocking=True)
        if training:
            optimizer.zero_grad(set_to_none=True)
        with torch.set_grad_enabled(training):
            logits = model(inputs)
            loss = criterion(logits, targets)
            if training:
                loss.backward()
                if gradient_clip_norm > 0:
                    torch.nn.utils.clip_grad_norm_(model.parameters(), gradient_clip_norm)
                optimizer.step()
        predictions = logits.argmax(dim=1)
        update_confusion(confusion, predictions, targets)
        total_loss += float(loss.detach())
        probabilities = logits.softmax(dim=1).detach().cpu()
        for index, (rows, columns) in enumerate(batch["shapes"]):
            predicted_tables = prediction_to_tables(
                predictions[index, :rows, :columns],
                probabilities[index, :, :rows, :columns],
                minimum_component_cells,
                header_fraction_threshold,
            )
            comparison = compare_table_sets(predicted_tables, batch["metadata"][index]["tables"])
            for key, value in comparison.items():
                box_totals[key] += value
            sample_count += 1
        progress.set_postfix(batch=f"{batch_index}/{len(loader)}", loss=f"{float(loss.detach()):.4f}")
    result = metrics_from_confusion(confusion)
    result["loss"] = total_loss / max(1, len(loader))
    for key, value in box_totals.items():
        result[key] = value / max(1, sample_count)
    return result


def extract_table_payload(worksheet: Any, table: dict[str, Any]) -> dict[str, Any]:
    return extract_table_payload_with_values(worksheet, table)


def _excel_color(value: Any) -> str | None:
    if value is None or getattr(value, "type", None) != "rgb":
        return None
    rgb = str(getattr(value, "rgb", "") or "")
    if len(rgb) == 8:
        rgb = rgb[2:]
    if len(rgb) != 6 or not re.fullmatch(r"[0-9a-fA-F]{6}", rgb):
        return None
    return f"#{rgb.upper()}"


def _column_width_px(worksheet: Any, column: int) -> int:
    width = worksheet.column_dimensions[get_column_letter(column)].width
    return max(56, min(420, round(float(width or 11.5) * 7 + 12)))


def _row_height_px(worksheet: Any, row: int, is_header: bool) -> int:
    height = worksheet.row_dimensions[row].height
    if height is None:
        return 40 if is_header else 28
    return max(24, min(180, round(float(height) * 96 / 72)))


def _header_paths(
    data: list[list[Any]],
    merges: list[dict[str, int]],
    header_rows: int,
    left: int,
) -> list[dict[str, Any]]:
    if not data:
        return []
    resolved = [list(row) for row in data[:header_rows]]
    for merged in merges:
        anchor_row = int(merged["row"])
        anchor_column = int(merged["col"])
        if anchor_row >= header_rows:
            continue
        anchor = resolved[anchor_row][anchor_column]
        last_row = min(header_rows, anchor_row + int(merged["rowspan"]))
        last_column = min(len(resolved[0]), anchor_column + int(merged["colspan"]))
        for row in range(anchor_row, last_row):
            for column in range(anchor_column, last_column):
                resolved[row][column] = anchor

    column_codes: list[str] = [""] * len(resolved[0])
    path_row_count = header_rows
    last_row_integers = [_integer(value) for value in resolved[-1]]
    code_candidates = [
        clean_text(value) if re.fullmatch(r"\(?-?\d+\)?", clean_text(value)) else ""
        for value in resolved[-1]
    ]
    negative_sequence = last_row_integers == list(range(-1, -len(last_row_integers) - 1, -1))
    mostly_codes = sum(bool(value) for value in code_candidates) >= max(
        2, math.ceil(len(code_candidates) * 0.70)
    )
    if negative_sequence or mostly_codes:
        column_codes = code_candidates
        path_row_count -= 1

    result: list[dict[str, Any]] = []
    for column in range(len(resolved[0])):
        path: list[str] = []
        for row in range(path_row_count):
            label = clean_text(resolved[row][column])
            if label and (not path or path[-1] != label):
                path.append(label)
        result.append({
            "columnIndex": column,
            "excelColumn": get_column_letter(left + column),
            "headerPath": path,
            "label": path[-1] if path else get_column_letter(left + column),
            "columnCode": column_codes[column],
        })
    return result


def extract_table_payload_with_values(
    worksheet: Any,
    table: dict[str, Any],
    value_worksheet: Any | None = None,
) -> dict[str, Any]:
    top, left = int(table["top"]), int(table["left"])
    bottom, right = int(table["bottom"]), int(table["right"])
    header_rows = max(1, min(bottom - top + 1, int(table["header_bottom"]) - top + 1))
    data: list[list[Any]] = []
    formula_cells: list[dict[str, Any]] = []
    cell_styles: list[dict[str, Any]] = []
    for row in range(top, bottom + 1):
        output_row: list[Any] = []
        for column in range(left, right + 1):
            cell = worksheet.cell(row, column)
            source_value = _cell_value(worksheet, row, column)
            output_value = source_value
            if getattr(cell, "data_type", "") == "f" and value_worksheet is not None:
                cached_value = _cell_value(value_worksheet, row, column)
                if cached_value is not None:
                    output_value = cached_value
                formula_cells.append({
                    "row": row - top,
                    "col": column - left,
                    "formula": json_safe(source_value),
                    "cachedValue": json_safe(cached_value),
                })
            output_row.append(json_safe(output_value))

            alignment = cell.alignment
            fill = _excel_color(getattr(cell.fill, "fgColor", None)) if cell.fill.fill_type else None
            font_color = _excel_color(getattr(cell.font, "color", None))
            is_header = row < top + header_rows
            style = {
                "row": row - top,
                "col": column - left,
                "isHeader": is_header,
                "bold": bool(cell.font.bold),
                "italic": bool(cell.font.italic),
                "horizontal": alignment.horizontal,
                "vertical": alignment.vertical,
                "wrapText": bool(alignment.wrap_text),
                "numberFormat": cell.number_format if cell.number_format != "General" else None,
                "fill": fill,
                "fontColor": font_color,
            }
            if is_header or any(value for key, value in style.items() if key not in {"row", "col", "isHeader"}):
                cell_styles.append(style)
        data.append(output_row)

    merges = []
    for merged_range in worksheet.merged_cells.ranges:
        merge_left, merge_top, merge_right, merge_bottom = range_boundaries(str(merged_range))
        clip_top, clip_left = max(top, merge_top), max(left, merge_left)
        clip_bottom, clip_right = min(bottom, merge_bottom), min(right, merge_right)
        if clip_top > clip_bottom or clip_left > clip_right:
            continue
        merge = {
            "row": clip_top - top,
            "col": clip_left - left,
            "rowspan": clip_bottom - clip_top + 1,
            "colspan": clip_right - clip_left + 1,
        }
        if merge["rowspan"] > 1 or merge["colspan"] > 1:
            merges.append(merge)
            if merge_top < top or merge_left < left:
                data[merge["row"]][merge["col"]] = json_safe(
                    _cell_value(worksheet, merge_top, merge_left)
                )

    sorted_merges = sorted(merges, key=lambda item: (item["row"], item["col"]))
    rows = bottom - top + 1
    columns = right - left + 1
    return {
        "sourceRange": f"{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}",
        "headerRows": header_rows,
        "rowCount": rows,
        "columnCount": columns,
        "valueRowCount": max(0, rows - header_rows),
        "data": data,
        "mergeCells": sorted_merges,
        "columns": _header_paths(data, sorted_merges, header_rows, left),
        "columnWidths": [_column_width_px(worksheet, column) for column in range(left, right + 1)],
        "rowHeights": [
            _row_height_px(worksheet, row, row < top + header_rows)
            for row in range(top, bottom + 1)
        ],
        "cellStyles": cell_styles,
        "formulaCells": formula_cells,
        "confidence": table.get("confidence"),
        "postprocessing": table.get("postprocessing"),
        "trimmedFooterRows": int(table.get("trimmed_footer_rows", 0)),
    }


def load_artifact(model_path: Path, device: torch.device) -> tuple[EFormTableNet, dict[str, Any]]:
    checkpoint = torch.load(model_path, map_location=device, weights_only=True)
    if checkpoint.get("format_version") != FORMAT_VERSION:
        raise ValueError(f"Unsupported artifact: {checkpoint.get('format_version')}")
    config = checkpoint["model_config"]
    model = EFormTableNet(config["input_channels"], config["base_channels"]).to(device)
    model.load_state_dict(checkpoint["state_dict"])
    model.eval()
    return model, checkpoint


def resolve_device(device_name: str | None = None) -> torch.device:
    requested = (device_name or "auto").strip().lower()
    if requested == "auto":
        return torch.device("cuda" if torch.cuda.is_available() else "cpu")
    if requested.startswith("cuda") and not torch.cuda.is_available():
        raise ValueError("CUDA was requested but is not available.")
    return torch.device(requested)


@torch.inference_mode()
def predict_worksheet(
    model: nn.Module,
    worksheet: Any,
    device: torch.device,
    minimum_component_cells: int = 4,
    header_fraction_threshold: float = 0.30,
) -> list[dict[str, Any]]:
    rows, columns = actual_bounds(worksheet)
    if rows == 0 or columns == 0:
        return []
    features = encode_sheet(worksheet, rows, columns).unsqueeze(0).to(device)
    logits = model(features)[0]
    probabilities = logits.softmax(dim=0).cpu()
    prediction = logits.argmax(dim=0).cpu()
    return prediction_to_tables(
        prediction,
        probabilities,
        minimum_component_cells,
        header_fraction_threshold,
    )


class EFormPredictor:
    """Reusable inference runtime. Load the artifact once, then process many workbooks."""

    def __init__(self, model_path: Path, device_name: str | None = None) -> None:
        self.model_path = Path(model_path).resolve()
        if not self.model_path.is_file():
            raise FileNotFoundError(f"Model artifact not found: {self.model_path}")
        self.device = resolve_device(device_name)
        self.model, self.checkpoint = load_artifact(self.model_path, self.device)
        self.inference_config = self.checkpoint.get("inference_config", {})

    @property
    def model_info(self) -> dict[str, Any]:
        return {
            "name": self.checkpoint["model_name"],
            "format": self.checkpoint["format_version"],
            "device": str(self.device),
            "artifact": self.model_path.name,
        }

    def predict(self, excel_path: Path, source_name: str | None = None) -> dict[str, Any]:
        excel_path = Path(excel_path)
        workbook = load_workbook(excel_path, read_only=False, data_only=False, keep_links=False)
        value_workbook = load_workbook(excel_path, read_only=False, data_only=True, keep_links=False)
        sheets: list[dict[str, Any]] = []
        try:
            for sheet_index, worksheet in enumerate(workbook.worksheets):
                value_worksheet = value_workbook.worksheets[sheet_index]
                regions = predict_worksheet(
                    self.model,
                    worksheet,
                    self.device,
                    int(self.inference_config.get("minimum_component_cells", 4)),
                    float(self.inference_config.get("header_fraction_threshold", 0.30)),
                )
                regions = [refine_table_bounds(worksheet, region) for region in regions]
                sheets.append({
                    "sheetIndex": sheet_index,
                    "sheetName": worksheet.title,
                    "sheetState": worksheet.sheet_state,
                    "tables": [
                        extract_table_payload_with_values(worksheet, region, value_worksheet)
                        for region in regions
                    ],
                })
        finally:
            value_workbook.close()
            workbook.close()

        tables = [table for sheet in sheets for table in sheet["tables"]]
        return {
            "schemaVersion": "1.0",
            "model": self.checkpoint["model_name"],
            "modelFormat": self.checkpoint["format_version"],
            "device": str(self.device),
            "sourceFile": source_name or excel_path.name,
            "summary": {
                "sheetCount": len(sheets),
                "tableCount": len(tables),
                "rowCount": sum(table["rowCount"] for table in tables),
                "valueRowCount": sum(table["valueRowCount"] for table in tables),
                "cellCount": sum(table["rowCount"] * table["columnCount"] for table in tables),
            },
            "sheets": sheets,
        }


def predict_excel(model_path: Path, excel_path: Path, device_name: str | None = None) -> dict[str, Any]:
    return EFormPredictor(model_path, device_name).predict(excel_path)


def _demo_records(records: list[dict[str, Any]], count: int, seed: int) -> list[dict[str, Any]]:
    rng = random.Random(seed)
    result: list[dict[str, Any]] = []
    desired = [("test_id", math.ceil(count / 2)), ("test_ood", count // 2)]
    for split, split_count in desired:
        candidates = [record for record in records if record["split"] == split and record["tables"]]
        result.extend(rng.sample(candidates, min(split_count, len(candidates))))
    return result


@torch.inference_mode()
def build_demos(
    model: nn.Module,
    module_root: Path,
    records: list[dict[str, Any]],
    device: torch.device,
    count: int,
    seed: int,
    minimum_component_cells: int = 4,
    header_fraction_threshold: float = 0.30,
) -> list[dict[str, Any]]:
    demos = []
    for record in _demo_records(records, count, seed):
        workbook = load_workbook(module_root / record["source_file"], read_only=False, data_only=False, keep_links=False)
        try:
            worksheet = workbook.worksheets[record["sheet_index"]]
            predicted = predict_worksheet(
                model,
                worksheet,
                device,
                minimum_component_cells,
                header_fraction_threshold,
            )
            expected = [dict(table) for table in record["tables"]]
            comparison = compare_table_sets(predicted, expected)
            demos.append({
                "source_file": record["source_file"],
                "sheet_name": worksheet.title,
                "split": record["split"],
                "expected_regions": expected,
                "predicted_regions": predicted,
                "comparison": comparison,
                "expected_output": [extract_table_payload(worksheet, table) for table in expected],
                "actual_output": [extract_table_payload(worksheet, table) for table in predicted],
            })
        finally:
            workbook.close()
    return demos


def save_curves(history: list[dict[str, Any]], path: Path) -> None:
    import matplotlib

    matplotlib.use("Agg", force=True)
    import matplotlib.pyplot as plt

    epochs = [item["epoch"] for item in history]
    figure, axes = plt.subplots(1, 3, figsize=(15, 4))
    axes[0].plot(epochs, [item["train"]["loss"] for item in history], label="train")
    axes[0].plot(epochs, [item["validation"]["loss"] for item in history], label="validation")
    axes[0].set_title("Loss")
    axes[1].plot(epochs, [item["validation"]["macro_f1"] for item in history], label="macro F1")
    axes[1].plot(epochs, [item["validation"]["foreground_iou"] for item in history], label="table IoU")
    axes[1].set_title("Validation cell metrics")
    axes[2].plot(epochs, [item["validation"]["box_iou"] for item in history], label="box IoU")
    axes[2].plot(epochs, [item["validation"]["exact_region"] for item in history], label="exact region")
    axes[2].set_title("Validation region metrics")
    for axis in axes:
        axis.set_xlabel("Epoch")
        axis.grid(alpha=0.25)
        axis.legend()
    figure.tight_layout()
    figure.savefig(path, dpi=150)
    plt.close(figure)


def _small_table(table: dict[str, Any], max_rows: int = 10, max_columns: int = 12) -> str:
    data = table.get("data", [])
    header_rows = int(table.get("headerRows", 0))
    body = []
    for row_index, row in enumerate(data[:max_rows]):
        cells = []
        tag = "th" if row_index < header_rows else "td"
        for value in row[:max_columns]:
            content = html.escape("" if value is None else str(value)).replace("\n", "<br>")
            cells.append(f"<{tag}>{content}</{tag}>")
        body.append("<tr>" + "".join(cells) + "</tr>")
    if not body:
        return "<div class='empty'>Không dự đoán được bảng.</div>"
    return "<table>" + "".join(body) + "</table>"


def save_demo_html(demos: list[dict[str, Any]], path: Path) -> None:
    sections = []
    for demo in demos:
        expected = demo["expected_output"][0] if demo["expected_output"] else {}
        actual = demo["actual_output"][0] if demo["actual_output"] else {}
        sections.append(
            "<section>"
            f"<h3>{html.escape(demo['split'])} · {html.escape(demo['source_file'])}</h3>"
            f"<p>Expected: {html.escape(str(demo['expected_regions']))}<br>"
            f"Actual: {html.escape(str(demo['predicted_regions']))}<br>"
            f"Metrics: {html.escape(str(demo['comparison']))}</p>"
            "<div class='pair'><div><h4>Kết quả mong đợi</h4>"
            + _small_table(expected)
            + "</div><div><h4>Kết quả thực tế</h4>"
            + _small_table(actual)
            + "</div></div></section>"
        )
    document = """<!doctype html><meta charset='utf-8'><style>
    body{font-family:Arial,sans-serif;color:#172033}section{border-top:2px solid #94a3b8;padding:12px 0}
    .pair{display:grid;grid-template-columns:1fr 1fr;gap:16px;overflow:auto}table{border-collapse:collapse;font-size:11px}
    th,td{border:1px solid #64748b;padding:4px;min-width:55px;max-width:160px}th{background:#dbeafe}.empty{color:#b91c1c}
    p{font-size:12px;white-space:normal}@media(max-width:900px){.pair{grid-template-columns:1fr}}</style>
    <h1>EForm: kết quả mong đợi và thực tế</h1>""" + "".join(sections)
    path.write_text(document, encoding="utf-8")


def _limit_records(records: list[dict[str, Any]], limit_per_split: int) -> list[dict[str, Any]]:
    if limit_per_split <= 0:
        return records
    grouped: defaultdict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        grouped[record["split"]].append(record)
    return [record for split in sorted(grouped) for record in grouped[split][:limit_per_split]]


def train_model(
    module_root: Path,
    output_dir: Path,
    epochs: int = 20,
    batch_size: int = 8,
    learning_rate: float = 1e-3,
    weight_decay: float = 1e-4,
    patience: int = 5,
    min_delta: float = 1e-5,
    gradient_clip_norm: float = 2.0,
    base_channels: int = 32,
    class_weights: tuple[float, float, float] = (0.25, 1.0, 1.15),
    device_name: str = "auto",
    num_workers: int = 0,
    minimum_component_cells: int = 4,
    header_fraction_threshold: float = 0.30,
    seed: int = 20260913,
    demo_samples: int = 6,
    limit_per_split: int = 0,
) -> dict[str, Any]:
    module_root = module_root.resolve()
    output_dir = output_dir.resolve()
    dataset_path = module_root / "Data" / "dataset.jsonl"
    dataset_report_path = module_root / "Data" / "dataset_report.json"
    if not dataset_path.is_file() or not dataset_report_path.is_file():
        raise FileNotFoundError("Thiếu Data/dataset.jsonl hoặc Data/dataset_report.json. Hãy chạy BuildDataset.ipynb trước.")
    dataset_report = json.loads(dataset_report_path.read_text(encoding="utf-8"))
    if not dataset_report.get("ready_for_training"):
        raise ValueError("dataset_report.json chưa đạt ready_for_training=true")
    all_records = [
        json.loads(line)
        for line in dataset_path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    records = _limit_records(all_records, limit_per_split)
    by_split = {
        split: [record for record in records if record["split"] == split]
        for split in ("train", "validation", "test_id", "test_ood")
    }
    if any(not by_split[split] for split in by_split):
        raise ValueError({split: len(items) for split, items in by_split.items()})
    if epochs < 1 or batch_size < 1 or base_channels < 1 or num_workers < 0:
        raise ValueError("epochs, batch_size, base_channels phải >= 1 và num_workers phải >= 0")
    if len(class_weights) != len(CLASS_NAMES):
        raise ValueError(f"class_weights phải có đúng {len(CLASS_NAMES)} giá trị")
    if minimum_component_cells < 1 or not 0.0 <= header_fraction_threshold <= 1.0:
        raise ValueError("minimum_component_cells phải >= 1 và header_fraction_threshold phải trong [0, 1]")

    seed_everything(seed)
    requested_device = device_name.strip().lower()
    device = torch.device(
        "cuda" if requested_device == "auto" and torch.cuda.is_available()
        else "cpu" if requested_device == "auto"
        else requested_device
    )
    if device.type == "cuda" and not torch.cuda.is_available():
        raise RuntimeError("DEVICE yêu cầu CUDA nhưng PyTorch không phát hiện GPU CUDA")
    if device.type == "cpu":
        torch.set_num_threads(max(1, min(8, torch.get_num_threads())))
    print("Device:", device)
    if device.type == "cuda":
        print("GPU:", torch.cuda.get_device_name(0))
    print("PyTorch:", torch.__version__)
    print("Samples:", {split: len(items) for split, items in by_split.items()})

    train_dataset = ExcelGridDataset(module_root, by_split["train"], include_augmentations=True)
    validation_dataset = ExcelGridDataset(module_root, by_split["validation"], include_augmentations=False)
    test_id_dataset = ExcelGridDataset(module_root, by_split["test_id"], include_augmentations=False)
    test_ood_dataset = ExcelGridDataset(module_root, by_split["test_ood"], include_augmentations=False)
    generator = torch.Generator().manual_seed(seed)
    loader_options = {
        "batch_size": batch_size,
        "num_workers": num_workers,
        "collate_fn": collate_grids,
        "pin_memory": device.type == "cuda",
        "persistent_workers": num_workers > 0,
    }
    train_loader = DataLoader(train_dataset, shuffle=True, generator=generator, **loader_options)
    validation_loader = DataLoader(validation_dataset, shuffle=False, **loader_options)
    test_id_loader = DataLoader(test_id_dataset, shuffle=False, **loader_options)
    test_ood_loader = DataLoader(test_ood_dataset, shuffle=False, **loader_options)

    model = EFormTableNet(base_channels=base_channels).to(device)
    class_weight_tensor = torch.tensor(class_weights, dtype=torch.float32, device=device)
    criterion = nn.CrossEntropyLoss(weight=class_weight_tensor, ignore_index=IGNORE_INDEX)
    optimizer = torch.optim.AdamW(model.parameters(), lr=learning_rate, weight_decay=weight_decay)
    scheduler = torch.optim.lr_scheduler.CosineAnnealingLR(optimizer, T_max=max(1, epochs))
    parameter_count = sum(parameter.numel() for parameter in model.parameters())
    print(f"Parameters: {parameter_count:,}")
    print(f"Batch size: {batch_size}; batches/epoch: {len(train_loader)}; epochs: {epochs}")
    epoch_options = {
        "gradient_clip_norm": gradient_clip_norm,
        "minimum_component_cells": minimum_component_cells,
        "header_fraction_threshold": header_fraction_threshold,
    }

    history: list[dict[str, Any]] = []
    best_state = copy.deepcopy(model.state_dict())
    best_score = -1.0
    best_epoch = 0
    stale_epochs = 0
    for epoch in range(1, epochs + 1):
        train_metrics = run_epoch(
            model,
            train_loader,
            criterion,
            device,
            optimizer,
            f"Epoch {epoch:02d}/{epochs} train",
            **epoch_options,
        )
        validation_metrics = run_epoch(
            model,
            validation_loader,
            criterion,
            device,
            None,
            f"Epoch {epoch:02d}/{epochs} val",
            **epoch_options,
        )
        scheduler.step()
        score = validation_metrics["mean_header_value_iou"]
        history.append({
            "epoch": epoch,
            "learning_rate": optimizer.param_groups[0]["lr"],
            "train": train_metrics,
            "validation": validation_metrics,
        })
        print(
            f"Epoch {epoch:02d}/{epochs} | "
            f"train_loss={train_metrics['loss']:.4f} | val_loss={validation_metrics['loss']:.4f} | "
            f"val_acc={validation_metrics['cell_accuracy']:.4f} | val_macro_f1={validation_metrics['macro_f1']:.4f} | "
            f"val_table_iou={validation_metrics['foreground_iou']:.4f} | val_box_iou={validation_metrics['box_iou']:.4f}"
        )
        if score > best_score + min_delta:
            best_score = score
            best_epoch = epoch
            best_state = copy.deepcopy(model.state_dict())
            stale_epochs = 0
        else:
            stale_epochs += 1
            if stale_epochs >= patience:
                print(f"Early stopping tại epoch {epoch}; best epoch={best_epoch}.")
                break

    model.load_state_dict(best_state)
    validation_metrics = run_epoch(
        model, validation_loader, criterion, device, None, "Best model validation", **epoch_options
    )
    test_id_metrics = run_epoch(
        model, test_id_loader, criterion, device, None, "Test-ID", **epoch_options
    )
    test_ood_metrics = run_epoch(
        model, test_ood_loader, criterion, device, None, "Test-OOD", **epoch_options
    )
    output_dir.mkdir(parents=True, exist_ok=True)
    model_path = output_dir / MODEL_FILENAME
    checkpoint = {
        "format_version": FORMAT_VERSION,
        "model_name": MODEL_NAME,
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "torch_version": str(torch.__version__),
        "model_config": {
            "input_channels": FEATURE_CHANNELS,
            "base_channels": base_channels,
            "classes": list(CLASS_NAMES),
        },
        "feature_config": {"channels": FEATURE_CHANNELS, "text_hash_dimensions": TEXT_HASH_DIM},
        "dataset_sha256": dataset_report["dataset_sha256"],
        "training_config": {
            "epochs_requested": epochs,
            "epochs_completed": len(history),
            "best_epoch": best_epoch,
            "batch_size": batch_size,
            "learning_rate": learning_rate,
            "weight_decay": weight_decay,
            "patience": patience,
            "min_delta": min_delta,
            "gradient_clip_norm": gradient_clip_norm,
            "class_weights": list(class_weights),
            "device": device_name,
            "num_workers": num_workers,
            "seed": seed,
            "limit_per_split": limit_per_split,
        },
        "inference_config": {
            "minimum_component_cells": minimum_component_cells,
            "header_fraction_threshold": header_fraction_threshold,
        },
        "metrics": {
            "validation": validation_metrics,
            "test_id": test_id_metrics,
            "test_ood": test_ood_metrics,
        },
        "state_dict": {key: value.detach().cpu() for key, value in best_state.items()},
    }
    torch.save(checkpoint, model_path)

    # Reload the single-file artifact before demos. This verifies the same file
    # that a future backend loader will receive.
    reloaded_model, _ = load_artifact(model_path, device)
    demos = build_demos(
        reloaded_model,
        module_root,
        all_records,
        device,
        demo_samples,
        seed + 1,
        minimum_component_cells,
        header_fraction_threshold,
    )
    report = {
        "model_name": MODEL_NAME,
        "model_file": str(model_path),
        "format_version": FORMAT_VERSION,
        "device": str(device),
        "torch_version": str(torch.__version__),
        "parameter_count": parameter_count,
        "dataset_counts": {split: len(items) for split, items in by_split.items()},
        "effective_train_samples": len(train_dataset),
        "training_config": checkpoint["training_config"],
        "history": history,
        "metrics": checkpoint["metrics"],
        "random_test_demos": demos,
        "artifact_reload_verified": True,
    }
    report_path = output_dir / "training_report.json"
    curves_path = output_dir / "training_curves.png"
    comparison_path = output_dir / "expected_vs_actual.html"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=json_safe), encoding="utf-8")
    save_curves(history, curves_path)
    save_demo_html(demos, comparison_path)
    print(json.dumps({
        "model_file": str(model_path),
        "best_epoch": best_epoch,
        "validation": validation_metrics,
        "test_id": test_id_metrics,
        "test_ood": test_ood_metrics,
        "artifact_reload_verified": True,
    }, ensure_ascii=False, indent=2))
    return report


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Train EForm Excel table detector")
    parser.add_argument("--module-root", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--epochs", type=int, default=20)
    parser.add_argument("--batch-size", type=int, default=8)
    parser.add_argument("--learning-rate", type=float, default=1e-3)
    parser.add_argument("--weight-decay", type=float, default=1e-4)
    parser.add_argument("--patience", type=int, default=5)
    parser.add_argument("--min-delta", type=float, default=1e-5)
    parser.add_argument("--gradient-clip-norm", type=float, default=2.0)
    parser.add_argument("--base-channels", type=int, default=32)
    parser.add_argument("--class-weights", type=float, nargs=3, default=(0.25, 1.0, 1.15))
    parser.add_argument("--device", default="auto")
    parser.add_argument("--num-workers", type=int, default=0)
    parser.add_argument("--minimum-component-cells", type=int, default=4)
    parser.add_argument("--header-fraction-threshold", type=float, default=0.30)
    parser.add_argument("--seed", type=int, default=20260913)
    parser.add_argument("--demo-samples", type=int, default=6)
    parser.add_argument("--limit-per-split", type=int, default=0)
    return parser.parse_args()


if __name__ == "__main__":
    arguments = parse_args()
    train_model(
        module_root=arguments.module_root,
        output_dir=arguments.output_dir,
        epochs=arguments.epochs,
        batch_size=arguments.batch_size,
        learning_rate=arguments.learning_rate,
        weight_decay=arguments.weight_decay,
        patience=arguments.patience,
        min_delta=arguments.min_delta,
        gradient_clip_norm=arguments.gradient_clip_norm,
        base_channels=arguments.base_channels,
        class_weights=tuple(arguments.class_weights),
        device_name=arguments.device,
        num_workers=arguments.num_workers,
        minimum_component_cells=arguments.minimum_component_cells,
        header_fraction_threshold=arguments.header_fraction_threshold,
        seed=arguments.seed,
        demo_samples=arguments.demo_samples,
        limit_per_split=arguments.limit_per_split,
    )
