from __future__ import annotations

import hashlib
import json
import math
import os
import re
import tempfile
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, time, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils.cell import range_boundaries


SCHEMA_VERSION = "excel-table-regions-1.0"
SUPPORTED_SUFFIXES = {".xlsx", ".xlsm"}
STANDARD_CODE_RE = re.compile(r"^(?:\(\s*\d+\s*\)|[A-Z]{1,3})$")
PAREN_CODE_RE = re.compile(r"^\(\s*\d+\s*\)$")
NUMBER_IN_TEXT_RE = re.compile(r"\d+(?:[.,]\d+)*")

FOOTER_TERMS = (
    "nguoi lap bieu",
    "nguoi kiem tra",
    "thu truong",
    "giam doc",
    "ky, ghi ro",
)
METADATA_TERMS = (
    "bieu so",
    "don vi bao cao",
    "don vi nhan bao cao",
    "don vi tinh",
    "ngay nhan bao cao",
)


@dataclass(frozen=True)
class BuildOptions:
    module_root: Path
    augment: bool = True
    minimum_family_samples: int = 12
    maximum_augmentations_per_sheet: int = 5
    include_hidden_sheets: bool = True
    max_rows: int = 50_000
    max_columns: int = 512
    max_cells: int = 2_000_000


def clean_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value if value is not None else "")).strip()


def fold_text(value: Any) -> str:
    text = unicodedata.normalize("NFD", clean_text(value).lower())
    text = "".join(character for character in text if unicodedata.category(character) != "Mn")
    return re.sub(r"\s+", " ", text.replace("đ", "d")).strip()


def json_safe(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else str(value)
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    return str(value)


def stable_hash(value: Any, length: int = 64) -> str:
    payload = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=json_safe)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()[:length]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _write_atomic(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8", newline="\n") as stream:
            stream.write(text)
        Path(temporary).replace(path)
    except Exception:
        Path(temporary).unlink(missing_ok=True)
        raise


def _actual_bounds(worksheet: Any) -> tuple[int, int]:
    max_row = 0
    max_column = 0
    for cell in getattr(worksheet, "_cells", {}).values():
        value = getattr(cell, "value", None)
        if value is not None and (not isinstance(value, str) or value.strip()):
            max_row = max(max_row, int(cell.row))
            max_column = max(max_column, int(cell.column))
    for merged_range in worksheet.merged_cells.ranges:
        _, _, right, bottom = range_boundaries(str(merged_range))
        max_row = max(max_row, bottom)
        max_column = max(max_column, right)
    return max_row, max_column


def _cell_value(worksheet: Any, row: int, column: int) -> Any:
    cell = worksheet.cell(row, column)
    return None if isinstance(cell, MergedCell) else cell.value


def _int_value(value: Any) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str) and re.fullmatch(r"[-+]?\d+", value.strip()):
        return int(value.strip())
    return None


def _row_summary(worksheet: Any, row: int, column_count: int) -> dict[str, Any]:
    non_empty: list[tuple[int, Any]] = []
    text_count = 0
    numeric_count = 0
    bold_count = 0
    centered_count = 0
    border_count = 0
    for column in range(1, column_count + 1):
        cell = worksheet.cell(row, column)
        value = _cell_value(worksheet, row, column)
        if value is None or (isinstance(value, str) and not value.strip()):
            continue
        non_empty.append((column, value))
        if isinstance(value, str) and not value.startswith("="):
            text_count += 1
        elif isinstance(value, (int, float)) and not isinstance(value, bool):
            numeric_count += 1
        bold_count += int(bool(getattr(cell.font, "bold", False)))
        centered_count += int(getattr(cell.alignment, "horizontal", None) in {"center", "centerContinuous"})
        border_count += sum(
            bool(getattr(getattr(cell.border, side, None), "style", None))
            for side in ("left", "right", "top", "bottom")
        )
    count = len(non_empty)
    texts = [fold_text(value) for _, value in non_empty]
    return {
        "row": row,
        "non_empty": non_empty,
        "count": count,
        "text_ratio": text_count / max(1, count),
        "numeric_ratio": numeric_count / max(1, count),
        "bold_ratio": bold_count / max(1, count),
        "center_ratio": centered_count / max(1, count),
        "border_density": border_count / max(1, count * 4),
        "folded_texts": texts,
        "joined": " | ".join(texts),
    }


def _is_blank(summary: dict[str, Any]) -> bool:
    return summary["count"] == 0


def _is_footer(summary: dict[str, Any]) -> bool:
    joined = summary["joined"]
    first = summary["folded_texts"][0].lstrip("* ") if summary["folded_texts"] else ""
    # "Ghi chú" is often a legitimate last header column.  It is a footer only
    # when it starts a sparse row below the data table.
    if summary["count"] <= 2 and first.startswith(("ghi chu", "chu thich")):
        return True
    return any(term in joined for term in FOOTER_TERMS)


def _is_metadata_or_title(summary: dict[str, Any]) -> bool:
    joined = summary["joined"]
    if any(term in joined for term in METADATA_TERMS):
        return True
    return (
        summary["count"] <= 3
        and (summary["bold_ratio"] >= 0.34 or summary["center_ratio"] >= 0.34)
        and any(text.startswith(("bao cao", "ket qua", "thong ke")) for text in summary["folded_texts"])
    )


def _numeric_sequence_columns(summary: dict[str, Any]) -> tuple[list[int], str] | None:
    """Recognize code rows such as -1..-13 and A,B,1..24.

    Alignment is checked against the physical column, so a deliberately blank
    code cell does not make a valid sequence fail (for example -6, blank, -8).
    """
    items = summary["non_empty"]
    numeric = [(column, _int_value(value)) for column, value in items if _int_value(value) is not None]
    textual = [
        (column, clean_text(value).upper())
        for column, value in items if _int_value(value) is None
    ]
    if len(numeric) < 2:
        return None
    if textual:
        textual_codes = [value for _, value in textual]
        expected_prefix = [chr(ord("A") + index) for index in range(len(textual_codes))]
        if textual_codes != expected_prefix:
            return None
    first_numeric_column = min(column for column, _ in numeric)
    if any(column > first_numeric_column for column, _ in textual):
        return None
    values = [value for _, value in numeric]
    if values[0] == -1 and all(value < 0 for value in values):
        offsets = {column + value for column, value in numeric}
        if len(offsets) == 1:
            return [column for column, _ in items], "negative_code_sequence"
    if values[0] == 1 and all(value > 0 for value in values):
        offsets = {column - value for column, value in numeric}
        if len(offsets) == 1:
            return [column for column, _ in items], "positive_code_sequence"
    return None


def _code_row_candidate(summary: dict[str, Any]) -> tuple[list[int], str, float] | None:
    standard = [
        column for column, value in summary["non_empty"]
        if STANDARD_CODE_RE.fullmatch(clean_text(value).upper())
    ]
    parenthetical_count = sum(
        bool(PAREN_CODE_RE.fullmatch(clean_text(value)))
        for _, value in summary["non_empty"]
    )
    if len(standard) >= 2 and parenthetical_count >= 1 and len(standard) / max(1, summary["count"]) >= 0.60:
        return standard, "standard_code_sequence", min(0.99, 0.86 + 0.02 * parenthetical_count)
    numeric_sequence = _numeric_sequence_columns(summary)
    if numeric_sequence:
        columns, method = numeric_sequence
        return columns, method, 0.98
    return None


def _header_top(summaries: dict[int, dict[str, Any]], code_row: int) -> int:
    top = code_row
    for row in range(code_row - 1, max(0, code_row - 13), -1):
        summary = summaries[row]
        if _is_blank(summary) or _is_footer(summary):
            break
        if row < code_row - 1 and _is_metadata_or_title(summary):
            break
        top = row
    return top


def _next_non_blank(summaries: dict[int, dict[str, Any]], start: int, end: int) -> int | None:
    for row in range(start, end + 1):
        if not _is_blank(summaries[row]):
            return row
    return None


def _data_bottom(
    summaries: dict[int, dict[str, Any]],
    code_row: int,
    row_count: int,
    left: int,
    right: int,
    next_table_top: int | None,
) -> int:
    hard_bottom = min(row_count, next_table_top - 1 if next_table_top else row_count)
    last = code_row
    pending_blank: int | None = None
    for row in range(code_row + 1, hard_bottom + 1):
        summary = summaries[row]
        if _is_footer(summary):
            break
        populated = [
            value for column, value in summary["non_empty"]
            if left <= column <= right
        ]
        if not populated:
            following = _next_non_blank(summaries, row + 1, min(hard_bottom, row + 1))
            if following is None or _is_footer(summaries[following]):
                break
            pending_blank = row
            continue
        if pending_blank is not None:
            last = row
            pending_blank = None
        else:
            last = row
    return last


def _fallback_table(
    summaries: dict[int, dict[str, Any]], row_count: int, column_count: int
) -> dict[str, Any] | None:
    data_start: int | None = None
    for row in range(2, row_count + 1):
        current = summaries[row]
        previous = summaries[row - 1]
        if (
            current["count"] >= 2
            and current["numeric_ratio"] >= 0.30
            and previous["count"] >= 2
            and previous["text_ratio"] >= 0.45
        ):
            data_start = row
            break
    if data_start is None:
        return None
    header_bottom = data_start - 1
    top = _header_top(summaries, data_start)
    # _header_top treats data_start like a code row; never include the data row itself.
    top = min(top, header_bottom)
    left = min(
        (column for row in range(top, data_start + 1) for column, _ in summaries[row]["non_empty"]),
        default=1,
    )
    right = max(
        (column for row in range(top, data_start + 1) for column, _ in summaries[row]["non_empty"]),
        default=column_count,
    )
    bottom = data_start
    for row in range(data_start + 1, row_count + 1):
        if _is_footer(summaries[row]) or _is_blank(summaries[row]):
            break
        if any(left <= column <= right for column, _ in summaries[row]["non_empty"]):
            bottom = row
    return {
        "top": top,
        "left": left,
        "bottom": bottom,
        "right": right,
        "header_bottom": header_bottom,
        "method": "fallback_numeric_transition",
        "confidence": 0.62,
    }


def detect_tables(worksheet: Any, row_count: int, column_count: int) -> tuple[list[dict[str, int]], str, float, list[str]]:
    if row_count == 0 or column_count == 0:
        return [], "empty_sheet", 1.0, []
    summaries = {row: _row_summary(worksheet, row, column_count) for row in range(1, row_count + 1)}
    candidates: list[tuple[int, list[int], str, float]] = []
    for row, summary in summaries.items():
        candidate = _code_row_candidate(summary)
        if candidate:
            columns, method, confidence = candidate
            candidates.append((row, columns, method, confidence))
    tops = [_header_top(summaries, row) for row, _, _, _ in candidates]
    tables: list[dict[str, int]] = []
    methods: list[str] = []
    confidences: list[float] = []
    warnings: list[str] = []
    for index, ((code_row, columns, method, confidence), top) in enumerate(zip(candidates, tops)):
        left, right = min(columns), max(columns)
        next_top = tops[index + 1] if index + 1 < len(tops) else None
        bottom = _data_bottom(summaries, code_row, row_count, left, right, next_top)
        if bottom == code_row:
            warnings.append("HEADER_ONLY_TABLE")
            confidence = min(confidence, 0.82)
        tables.append({
            "top": top,
            "left": left,
            "bottom": bottom,
            "right": right,
            "header_bottom": code_row,
        })
        methods.append(method)
        confidences.append(confidence)
    if not tables:
        fallback = _fallback_table(summaries, row_count, column_count)
        if fallback:
            tables.append({key: fallback[key] for key in ("top", "left", "bottom", "right", "header_bottom")})
            methods.append(str(fallback["method"]))
            confidences.append(float(fallback["confidence"]))
            warnings.append("FALLBACK_BOUNDARY")
    method = "+".join(sorted(set(methods))) if methods else "no_table"
    return tables, method, min(confidences, default=0.70), sorted(set(warnings))


def _header_signature(worksheet: Any, tables: list[dict[str, int]]) -> list[list[Any]]:
    result: list[list[Any]] = []
    for table in tables:
        cells: list[Any] = []
        for row in range(table["top"], table["header_bottom"] + 1):
            for column in range(table["left"], table["right"] + 1):
                value = _cell_value(worksheet, row, column)
                text = fold_text(value)
                if text:
                    text = NUMBER_IN_TEXT_RE.sub("<N>", text)[:200]
                cell = worksheet.cell(row, column)
                cells.append([
                    row - table["top"],
                    column - table["left"],
                    text,
                    bool(getattr(cell.font, "bold", False)),
                    getattr(cell.alignment, "horizontal", None) or "",
                    sum(bool(getattr(getattr(cell.border, side, None), "style", None)) for side in ("left", "right", "top", "bottom")),
                ])
        result.append(cells)
    return result


def _analyse_workbook(path: Path, options: BuildOptions) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    sheets: list[dict[str, Any]] = []
    signature_sheets: list[dict[str, Any]] = []
    try:
        for sheet_index, worksheet in enumerate(workbook.worksheets):
            if worksheet.sheet_state != "visible" and not options.include_hidden_sheets:
                continue
            row_count, column_count = _actual_bounds(worksheet)
            if row_count > options.max_rows or column_count > options.max_columns or row_count * column_count > options.max_cells:
                raise ValueError(f"Sheet {worksheet.title!r} vượt giới hạn {row_count}x{column_count}.")
            tables, method, confidence, warnings = detect_tables(worksheet, row_count, column_count)
            merges = sorted(str(item) for item in worksheet.merged_cells.ranges)
            sheets.append({
                "sheet_index": sheet_index,
                "sheet_name": worksheet.title,
                "sheet_shape": [row_count, column_count],
                "tables": tables,
                "annotation": {
                    "method": method,
                    "confidence": round(confidence, 4),
                    "warnings": warnings,
                },
            })
            signature_sheets.append({
                "sheet_index": sheet_index,
                "shape": [row_count, column_count],
                "merges": merges,
                "tables": tables,
                "headers": _header_signature(worksheet, tables),
            })
    finally:
        workbook.close()
    signature = {"sheet_count": len(sheets), "sheets": signature_sheets}
    return {"sheets": sheets, "layout_group": f"layout-{stable_hash(signature, 20)}"}


def _validate_tables(record: dict[str, Any]) -> list[str]:
    errors: list[str] = []
    rows, columns = record["sheet_shape"]
    for index, table in enumerate(record["tables"], 1):
        if not (
            1 <= table["top"] <= table["header_bottom"] <= table["bottom"] <= rows
            and 1 <= table["left"] <= table["right"] <= columns
        ):
            errors.append(f"table {index} vượt sheet hoặc sai thứ tự")
    return errors


def _split_workbooks(records: list[dict[str, Any]]) -> None:
    workbooks_by_family: defaultdict[str, set[str]] = defaultdict(set)
    for record in records:
        workbooks_by_family[record["layout_group"]].add(record["workbook_id"])
    split_by_workbook: dict[str, str] = {}
    for family, workbook_ids in sorted(workbooks_by_family.items()):
        ordered = sorted(workbook_ids)
        if len(ordered) <= 2 and int(stable_hash([family, "ood"], 8), 16) % 100 < 15:
            for workbook_id in ordered:
                split_by_workbook[workbook_id] = "test_ood"
            continue
        for workbook_id in ordered:
            bucket = int(stable_hash([workbook_id, "id-split"], 8), 16) % 100
            split_by_workbook[workbook_id] = "validation" if bucket < 10 else "test_id" if bucket < 25 else "train"
        if not any(split_by_workbook[item] == "train" for item in ordered):
            split_by_workbook[ordered[0]] = "train"
    for record in records:
        record["split"] = split_by_workbook[record["workbook_id"]]


def _augmentation_recipes(count: int, seed: str) -> list[dict[str, Any]]:
    catalog = [
        {"row_shift": 1, "column_shift": 0, "style_dropout": 0.10, "header_text_dropout": 0.00},
        {"row_shift": 0, "column_shift": 1, "style_dropout": 0.00, "header_text_dropout": 0.05},
        {"row_shift": 2, "column_shift": 1, "style_dropout": 0.15, "header_text_dropout": 0.00},
        {"row_shift": 4, "column_shift": 0, "style_dropout": 0.25, "header_text_dropout": 0.08},
        {"row_shift": 1, "column_shift": 2, "style_dropout": 0.05, "header_text_dropout": 0.12},
    ]
    offset = int(stable_hash(seed, 8), 16) % len(catalog)
    return [dict(catalog[(offset + index) % len(catalog)], seed=int(stable_hash([seed, index], 8), 16)) for index in range(count)]


def _attach_augmentations(records: list[dict[str, Any]], options: BuildOptions) -> None:
    train_by_family: defaultdict[str, list[dict[str, Any]]] = defaultdict(list)
    family_workbooks: defaultdict[str, set[str]] = defaultdict(set)
    for record in records:
        family_workbooks[record["layout_group"]].add(record["workbook_id"])
        if record["split"] == "train":
            train_by_family[record["layout_group"]].append(record)
    for record in records:
        record["augmentation_recipes"] = []
    if not options.augment:
        return
    for family, train_records in train_by_family.items():
        family_size = len(family_workbooks[family])
        missing = max(0, options.minimum_family_samples - family_size)
        if not train_records or not missing:
            continue
        per_record = min(options.maximum_augmentations_per_sheet, math.ceil(missing / len(train_records)))
        for record in train_records:
            record["augmentation_recipes"] = _augmentation_recipes(per_record, record["sample_id"])


def _apply_curated_labels(records: list[dict[str, Any]], curated_labels: list[dict[str, Any]]) -> int:
    by_key = {(item["source_file"], int(item["sheet_index"])): item for item in curated_labels}
    applied = 0
    for record in records:
        curated = by_key.get((record["source_file"], record["sheet_index"]))
        if not curated:
            continue
        original = record["tables"]
        record["tables"] = curated["tables"]
        errors = _validate_tables(record)
        if errors:
            record["tables"] = original
            raise ValueError(f"Curated label sai cho {record['source_file']}: {errors}")
        record["annotation"] = {
            "method": "curated_exact_output_alignment",
            "confidence": 1.0,
            "warnings": [],
        }
        applied += 1
    return applied


def build_dataset(
    options: BuildOptions | Path,
    *,
    curated_labels: list[dict[str, Any]] | None = None,
    progress: bool = True,
) -> dict[str, Any]:
    if isinstance(options, Path):
        options = BuildOptions(options)
    module_root = options.module_root.resolve()
    raw_root = module_root / "Data" / "Raw"
    if not raw_root.is_dir():
        raise FileNotFoundError(f"Không tìm thấy {raw_root}")
    files = sorted(
        path for path in raw_root.rglob("*")
        if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES and not path.name.startswith("~$")
    )
    digests: dict[Path, str] = {}
    canonical_by_digest: dict[str, Path] = {}
    for index, path in enumerate(files, 1):
        digest = sha256_file(path)
        digests[path] = digest
        canonical_by_digest.setdefault(digest, path)
        if progress and (index == 1 or index % 200 == 0 or index == len(files)):
            print(f"Đã kiểm tra SHA-256: {index}/{len(files)}")

    records: list[dict[str, Any]] = []
    parse_errors: list[dict[str, str]] = []
    canonical_paths = sorted(canonical_by_digest.values())
    for index, path in enumerate(canonical_paths, 1):
        digest = digests[path]
        source_file = path.relative_to(module_root).as_posix()
        workbook_id = f"wb-{digest[:24]}"
        try:
            analysis = _analyse_workbook(path, options)
            for sheet in analysis["sheets"]:
                sample_id = f"sheet-{stable_hash([workbook_id, sheet['sheet_index']], 24)}"
                record = {
                    "schema_version": SCHEMA_VERSION,
                    "sample_id": sample_id,
                    "workbook_id": workbook_id,
                    "source_file": source_file,
                    "source_sha256": digest,
                    "sheet_index": sheet["sheet_index"],
                    "sheet_name": sheet["sheet_name"],
                    "sheet_shape": sheet["sheet_shape"],
                    "layout_group": analysis["layout_group"],
                    "tables": sheet["tables"],
                    "annotation": sheet["annotation"],
                }
                errors = _validate_tables(record)
                if errors:
                    raise ValueError(f"Nhãn không hợp lệ: {errors}")
                records.append(record)
        except Exception as exc:
            parse_errors.append({"source_file": source_file, "error_type": type(exc).__name__, "error": str(exc)})
        if progress and (index == 1 or index % 100 == 0 or index == len(canonical_paths)):
            print(f"Đã tạo nhãn: {index}/{len(canonical_paths)} workbook duy nhất")

    curated_count = _apply_curated_labels(records, curated_labels or [])
    _split_workbooks(records)
    _attach_augmentations(records, options)
    records.sort(key=lambda item: (item["source_file"].lower(), item["sheet_index"]))

    sample_ids = [record["sample_id"] for record in records]
    workbook_splits: defaultdict[str, set[str]] = defaultdict(set)
    family_splits: defaultdict[str, set[str]] = defaultdict(set)
    for record in records:
        workbook_splits[record["workbook_id"]].add(record["split"])
        family_splits[record["layout_group"]].add(record["split"])
    invalid_labels = [
        {"sample_id": record["sample_id"], "errors": _validate_tables(record)}
        for record in records if _validate_tables(record)
    ]
    overlapping_tables = []
    for record in records:
        for first_index, first in enumerate(record["tables"]):
            for second in record["tables"][first_index + 1:]:
                rows_overlap = first["top"] <= second["bottom"] and second["top"] <= first["bottom"]
                columns_overlap = first["left"] <= second["right"] and second["left"] <= first["right"]
                if rows_overlap and columns_overlap:
                    overlapping_tables.append(record["sample_id"])
                    break
    # A family may intentionally be present in Train and Test-ID.  Test-OOD is
    # the only split that must have no sibling from the same family elsewhere.
    ood_leakage = sorted(
        family for family, splits in family_splits.items()
        if "test_ood" in splits and len(splits) > 1
    )
    validation = {
        "sample_ids_unique": len(sample_ids) == len(set(sample_ids)),
        "one_split_per_workbook": all(len(splits) == 1 for splits in workbook_splits.values()),
        "test_ood_has_no_family_leakage": not ood_leakage,
        "test_ood_leaking_families": ood_leakage,
        "table_labels_valid": not invalid_labels,
        "invalid_label_examples": invalid_labels[:10],
        "table_regions_do_not_overlap": not overlapping_tables,
        "overlapping_table_examples": overlapping_tables[:10],
        "augmentation_only_on_train": all(
            record["split"] == "train" or not record["augmentation_recipes"]
            for record in records
        ),
        "augmentation_parameters_valid": all(
            recipe["row_shift"] >= 0
            and recipe["column_shift"] >= 0
            and 0 <= recipe["style_dropout"] <= 1
            and 0 <= recipe["header_text_dropout"] <= 1
            for record in records for recipe in record["augmentation_recipes"]
        ),
        "contains_no_doctype_field": all("doc_type_code" not in record and "doctype" not in record for record in records),
        "contains_no_raw_cell_values": all("cells" not in record and "values" not in record for record in records),
    }
    validation["valid"] = all(
        value for key, value in validation.items()
        if key not in {
            "test_ood_leaking_families",
            "invalid_label_examples",
            "overlapping_table_examples",
        }
    )
    if not validation["valid"]:
        raise ValueError(f"Dataset không qua kiểm tra: {validation}")

    jsonl = "".join(json.dumps(record, ensure_ascii=False, separators=(",", ":")) + "\n" for record in records)
    dataset_path = module_root / "Data" / "dataset.jsonl"
    _write_atomic(dataset_path, jsonl)
    family_sizes = Counter(record["layout_group"] for record in records)
    split_counts = Counter(record["split"] for record in records)
    method_counts = Counter(record["annotation"]["method"] for record in records)
    warning_counts = Counter(warning for record in records for warning in record["annotation"]["warnings"])
    report = {
        "schema_version": SCHEMA_VERSION,
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "dataset_file": "Data/dataset.jsonl",
        "raw_workbook_count": len(files),
        "unique_workbook_count": len(canonical_paths),
        "duplicate_workbook_count": len(files) - len(canonical_paths),
        "parse_error_count": len(parse_errors),
        "parse_errors": parse_errors,
        "sheet_sample_count": len(records),
        "table_count": sum(len(record["tables"]) for record in records),
        "no_table_sheet_count": sum(not record["tables"] for record in records),
        "multi_table_sheet_count": sum(len(record["tables"]) > 1 for record in records),
        "layout_group_count": len(family_sizes),
        "singleton_layout_group_count": sum(size == 1 for size in family_sizes.values()),
        "split_sheet_counts": dict(sorted(split_counts.items())),
        "annotation_method_counts": dict(sorted(method_counts.items())),
        "annotation_warning_counts": dict(sorted(warning_counts.items())),
        "curated_label_count": curated_count,
        "augmentation_recipe_count": sum(len(record["augmentation_recipes"]) for record in records),
        "effective_training_sample_count": sum(
            1 + len(record["augmentation_recipes"])
            for record in records if record["split"] == "train"
        ),
        "validation": validation,
        "dataset_sha256": hashlib.sha256(jsonl.encode("utf-8")).hexdigest(),
        "ready_for_training": not parse_errors and validation["valid"],
        "important_note": "Annotation ngoài curated case được tạo tự động từ cấu trúc Excel; cần đánh giá Test-ID và Test-OOD trước khi công bố độ chính xác.",
    }
    _write_atomic(
        module_root / "Data" / "dataset_report.json",
        json.dumps(report, ensure_ascii=False, indent=2) + "\n",
    )
    if progress:
        print(
            f"Hoàn tất: {len(records)} sheet, {report['table_count']} bảng, "
            f"{report['augmentation_recipe_count']} augmentation recipes."
        )
    return report


def _value_type(value: Any, data_type: str) -> str:
    if value is None or (isinstance(value, str) and not value.strip()):
        return "blank"
    if data_type == "f" or (isinstance(value, str) and value.startswith("=")):
        return "formula"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, datetime):
        return "datetime"
    if isinstance(value, date):
        return "date"
    if isinstance(value, time):
        return "time"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, float):
        return "decimal"
    return "text"


def encode_record(
    module_root: Path,
    record: dict[str, Any],
    augmentation_index: int | None = None,
) -> tuple[dict[str, Any], list[dict[str, int]]]:
    """Materialize model features on demand; never include source raw values."""
    recipe = None if augmentation_index is None else record["augmentation_recipes"][augmentation_index]
    row_shift = int(recipe.get("row_shift", 0)) if recipe else 0
    column_shift = int(recipe.get("column_shift", 0)) if recipe else 0
    style_dropout = float(recipe.get("style_dropout", 0)) if recipe else 0.0
    header_dropout = float(recipe.get("header_text_dropout", 0)) if recipe else 0.0
    seed = int(recipe.get("seed", 0)) if recipe else 0
    workbook = load_workbook(module_root / record["source_file"], read_only=False, data_only=False, keep_links=False)
    try:
        worksheet = workbook.worksheets[record["sheet_index"]]
        rows, columns = record["sheet_shape"]
        header_cells = {
            (row, column)
            for table in record["tables"]
            for row in range(table["top"], table["header_bottom"] + 1)
            for column in range(table["left"], table["right"] + 1)
        }
        cells: list[dict[str, Any]] = []
        for row in range(1, rows + 1):
            row_numbers = [_int_value(_cell_value(worksheet, row, column)) for column in range(1, columns + 1)]
            numeric_values = [value for value in row_numbers if value is not None]
            negative_sequence = len(numeric_values) >= 2 and numeric_values == list(
                range(-1, -len(numeric_values) - 1, -1)
            )
            for column in range(1, columns + 1):
                cell = worksheet.cell(row, column)
                value = _cell_value(worksheet, row, column)
                value_type = _value_type(value, getattr(cell, "data_type", "") or "")
                if value_type == "blank" and not cell.has_style:
                    continue
                if value_type == "text":
                    token = NUMBER_IN_TEXT_RE.sub("<N>", fold_text(value))[:256]
                else:
                    token = f"<{value_type.upper()}>"
                drop_text_score = int(stable_hash([seed, row, column, "text"], 8), 16) / 0xFFFFFFFF
                if (row, column) in header_cells and drop_text_score < header_dropout:
                    token = "<MASK>"
                drop_style_score = int(stable_hash([seed, row, column, "style"], 8), 16) / 0xFFFFFFFF
                keep_style = drop_style_score >= style_dropout
                cells.append({
                    "row": row + row_shift,
                    "column": column + column_shift,
                    "token": token,
                    "value_type": value_type,
                    "is_negative": bool((_int_value(value) or 0) < 0),
                    "is_negative_sequence_row": negative_sequence,
                    "bold": bool(getattr(cell.font, "bold", False)) if keep_style else False,
                    "centered": getattr(cell.alignment, "horizontal", None) in {"center", "centerContinuous"} if keep_style else False,
                    "border_count": sum(
                        bool(getattr(getattr(cell.border, side, None), "style", None))
                        for side in ("left", "right", "top", "bottom")
                    ) if keep_style else 0,
                })
        merges = []
        for merged_range in sorted(str(item) for item in worksheet.merged_cells.ranges):
            left, top, right, bottom = range_boundaries(merged_range)
            merges.append({
                "top": top + row_shift,
                "left": left + column_shift,
                "bottom": bottom + row_shift,
                "right": right + column_shift,
            })
    finally:
        workbook.close()
    shifted_tables = [
        {
            "top": table["top"] + row_shift,
            "left": table["left"] + column_shift,
            "bottom": table["bottom"] + row_shift,
            "right": table["right"] + column_shift,
            "header_bottom": table["header_bottom"] + row_shift,
        }
        for table in record["tables"]
    ]
    return {
        "shape": [rows + row_shift, columns + column_shift],
        "cells": cells,
        "merges": merges,
    }, shifted_tables
