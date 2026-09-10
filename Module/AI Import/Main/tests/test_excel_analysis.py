import sys
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path

from openpyxl import Workbook


sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from ai_import.config import settings
from ai_import.excel import WorkbookAnalyzer


class ExcelAnalysisTests(unittest.TestCase):
    def test_separates_two_tables_and_never_puts_values_in_header_path(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "two-tables.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = "Biểu số: 01d/TP/TĐ-BH-TĐG"
            sheet.merge_cells("A2:C2")
            sheet["A2"] = "I. KẾT QUẢ THẨM ĐỊNH"
            sheet["A3"], sheet["B3"], sheet["C3"] = "Chỉ tiêu", "Tổng số", "Ghi chú"
            sheet["A4"], sheet["B4"], sheet["C4"] = "A", "(1)", "(2)"
            sheet["A5"], sheet["B5"], sheet["C5"] = "Dòng dữ liệu", 77, "hợp lệ"
            sheet.merge_cells("A7:C7")
            sheet["A7"] = "II. KẾT QUẢ BAN HÀNH"
            sheet["A8"], sheet["B8"], sheet["C8"] = "Chi tiết", "Tổng số", "Trong đó"
            sheet["A9"], sheet["B9"], sheet["C9"] = "A", "(1)", "(2)"
            sheet["A10"], sheet["B10"], sheet["C10"] = "Tổng số văn bản", 10, 8
            sheet["A11"], sheet["B11"], sheet["C11"] = "Trong đó: văn bản cấp xã", 4, 3
            workbook.save(path)
            workbook.close()

            analyzer = WorkbookAnalyzer(replace(settings, enforce_upload_root=False))
            analysis = analyzer.analyze(str(path), preview_rows=20)
            regions = analysis.sheets[0].regions
            self.assertEqual(len(regions), 2)
            self.assertEqual(regions[0].rows[0][1], 77)
            all_headers = " | ".join(part for region in regions for path_parts in region.header_paths for part in path_parts)
            self.assertNotIn("77", all_headers)
            self.assertEqual([item.level for item in regions[1].indicators], [0, 1])


if __name__ == "__main__":
    unittest.main()
