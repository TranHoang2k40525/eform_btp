from __future__ import annotations

import time
import re
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from .config import Settings
from .detection import choose_header_rows, find_row_bands, flatten_headers
from .hierarchy import enrich_indicators
from .models import IndicatorDto, SheetAnalysisDto, TableRegionDto, WorkbookAnalysisDto
from .security import validate_workbook_path
from .text import fold_vietnamese


_COLUMN_CODE_RE = re.compile(r"^(?:\(\s*\d+\s*\)|[A-Z]{1,3})$")
_NUMERIC_CODE_RE = re.compile(r"^\(\s*\d+\s*\)$")
_INDICATOR_CODE_RE = re.compile(r"^(?:[IVXLCDM]+\.?|\d+(?:\.\d+)*|[-–—])$", re.IGNORECASE)
_STOP_ROW_PREFIXES = ("ghi chu", "cot ", "nguoi lap bieu", "nguoi kiem tra", "thu truong don vi")


def _trim_matrix(matrix: list[list[Any]]) -> tuple[list[list[Any]], int, int]:
    non_empty = [
        (r, c)
        for r, row in enumerate(matrix)
        for c, value in enumerate(row)
        if value is not None and str(value).strip() != ""
    ]
    if not non_empty:
        return [], 0, 0
    max_row = max(item[0] for item in non_empty)
    max_col = max(item[1] for item in non_empty)
    return [row[: max_col + 1] for row in matrix[: max_row + 1]], max_row + 1, max_col + 1


def _read_merge_ranges(workbook_path: Path, worksheet_path: str) -> list[str]:
    """Read merge refs directly; read-only openpyxl skips drawings/images and merged-cell objects."""
    ranges: list[str] = []
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}mergeCell"
    normalized = worksheet_path.lstrip("/")
    with zipfile.ZipFile(workbook_path, "r") as archive, archive.open(normalized, "r") as stream:
        for _, element in ET.iterparse(stream, events=("end",)):
            if element.tag == namespace and element.attrib.get("ref"):
                ranges.append(element.attrib["ref"])
            element.clear()
    return ranges


def _text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _expanded_matrix(matrix: list[list[Any]], ranges: list[str]) -> list[list[Any]]:
    result = [list(row) for row in matrix]
    for ref in ranges:
        min_col, min_row, max_col, max_row = range_boundaries(ref)
        anchor = result[min_row - 1][min_col - 1] if min_row <= len(result) and min_col <= len(result[min_row - 1]) else None
        for row in range(min_row - 1, min(max_row, len(result))):
            for column in range(min_col - 1, max_col):
                if column < len(result[row]) and result[row][column] in (None, ""):
                    result[row][column] = anchor
    return result


def _code_rows(matrix: list[list[Any]]) -> list[tuple[int, list[int]]]:
    result: list[tuple[int, list[int]]] = []
    for row_index, row in enumerate(matrix):
        columns = [index for index, value in enumerate(row) if _COLUMN_CODE_RE.fullmatch(_text(value))]
        numeric = sum(bool(_NUMERIC_CODE_RE.fullmatch(_text(row[index]))) for index in columns)
        if numeric >= 2:
            result.append((row_index, columns))
    return result


def _header_start(matrix: list[list[Any]], code_row: int) -> int:
    lower_bound = max(0, code_row - 8)
    for row in range(code_row - 1, lower_bound - 1, -1):
        if not any(_text(value) for value in matrix[row]):
            return row + 1
    return lower_bound


def _is_metadata_header(value: str) -> bool:
    folded = fold_vietnamese(value)
    return (
        folded.startswith("bieu so")
        or folded.startswith("don vi bao cao")
        or folded.startswith("don vi nhan bao cao")
        or folded.startswith("ky bao cao")
        or folded.startswith("ngay nhan bao cao")
        or (len(value) > 180 and "bao cao" in folded)
    )


def _header_context(
    expanded: list[list[Any]], header_start: int, code_row: int, columns: list[int]
) -> tuple[list[list[str]], list[str], list[str]]:
    paths: list[list[str]] = []
    section_path: list[str] = []
    for row_index in range(header_start, code_row):
        unique = list(dict.fromkeys(_text(value) for value in expanded[row_index] if _text(value)))
        if len(unique) == 1 and re.match(r"^[IVXLCDM]+\s*[.)]", unique[0], re.IGNORECASE):
            section_path.append(unique[0])
    for column in columns:
        parts: list[str] = []
        for row_index in range(header_start, code_row):
            value = _text(expanded[row_index][column] if column < len(expanded[row_index]) else None)
            if not value or _COLUMN_CODE_RE.fullmatch(value) or _is_metadata_header(value):
                continue
            if not parts or fold_vietnamese(parts[-1]) != fold_vietnamese(value):
                parts.append(value)
        if section_path and (not parts or fold_vietnamese(parts[0]) != fold_vietnamese(section_path[-1])):
            parts = [section_path[-1], *parts]
        paths.append(parts)
    headers = [path[-1] if path else f"Cột {column + 1}" for path, column in zip(paths, columns)]
    return paths, headers, section_path


def _data_end(matrix: list[list[Any]], code_row: int, columns: list[int]) -> int:
    end = code_row
    for row_index in range(code_row + 1, len(matrix)):
        values = [_text(matrix[row_index][column] if column < len(matrix[row_index]) else None) for column in columns]
        non_empty = [value for value in values if value]
        if not non_empty:
            break
        first = fold_vietnamese(non_empty[0]).lstrip("* ")
        if len(non_empty) <= 2 and any(first.startswith(prefix) for prefix in _STOP_ROW_PREFIXES):
            break
        end = row_index
    return end


def _extract_indicators(
    rows: list[list[Any]], headers: list[str], codes: list[str], data_start_row: int
) -> list[IndicatorDto]:
    if len(rows) < 2:
        return []
    marker_column: int | None = None
    for column in range(min(3, len(headers))):
        raw_values = [row[column] for row in rows if column < len(row) and _text(row[column])]
        values = [_text(value) for value in raw_values if isinstance(value, str)]
        header = fold_vietnamese(headers[column])
        if "stt" in header or "so thu tu" in header:
            marker_column = column
            break
        if len(values) >= 2 and sum(bool(_INDICATOR_CODE_RE.fullmatch(value)) for value in values) / len(values) >= 0.70:
            marker_column = column
            break

    label_column: int | None = None
    label_terms = ("chi tiet", "noi dung", "ten don vi", "phan theo", "chi tieu")
    for column, header in enumerate(headers[:3]):
        if column != marker_column and any(term in fold_vietnamese(header) for term in label_terms):
            label_column = column
            break
    if label_column is None and marker_column is not None and marker_column + 1 < len(headers):
        candidate = marker_column + 1
        values = [_text(row[candidate]) for row in rows if candidate < len(row) and _text(row[candidate])]
        if values and sum(any(char.isalpha() for char in value) for value in values) / len(values) >= 0.70:
            label_column = candidate
    if label_column is None and codes and codes[0] in {"A", "B"}:
        values = [_text(row[0]) for row in rows if row and _text(row[0])]
        if len(rows) > 1 and values and sum(any(char.isalpha() for char in value) for value in values) / len(values) >= 0.70:
            label_column = 0
    if label_column is None:
        return []

    indicators = []
    for offset, row in enumerate(rows):
        label = _text(row[label_column] if label_column < len(row) else None)
        code = _text(row[marker_column] if marker_column is not None and marker_column < len(row) else None)
        if label:
            indicators.append(IndicatorDto(
                source_ref=f"R{data_start_row + offset}",
                source_row=data_start_row + offset,
                code=code,
                label=label,
            ))
    return enrich_indicators(indicators)


def _statistical_regions(
    sheet_name: str,
    matrix: list[list[Any]],
    merge_ranges: list[str],
    preview_rows: int,
) -> list[TableRegionDto]:
    expanded = _expanded_matrix(matrix, merge_ranges)
    regions: list[TableRegionDto] = []
    for code_row, columns in _code_rows(matrix):
        start = _header_start(matrix, code_row)
        end = _data_end(matrix, code_row, columns)
        if end <= code_row:
            continue
        paths, headers, section_path = _header_context(expanded, start, code_row, columns)
        codes = [_text(matrix[code_row][column]) for column in columns]
        raw_rows = [
            [matrix[row][column] if column < len(matrix[row]) else None for column in columns]
            for row in range(code_row + 1, min(end + 1, code_row + 1 + preview_rows))
        ]
        data_start = code_row + 2
        regions.append(TableRegionDto(
            sheet=sheet_name,
            start_row=start + 1,
            end_row=end + 1,
            start_column=min(columns) + 1,
            end_column=max(columns) + 1,
            header_start_row=start + 1,
            header_end_row=code_row + 1,
            data_start_row=data_start,
            confidence=min(1.0, 0.72 + 0.02 * len(columns)),
            headers=headers,
            source_columns=[column + 1 for column in columns],
            header_paths=paths,
            column_codes=codes,
            section_path=section_path,
            indicators=_extract_indicators(raw_rows, headers, codes, data_start),
            rows=raw_rows,
        ))
    return regions


class WorkbookAnalyzer:
    def __init__(self, config: Settings):
        self.config = config

    def analyze(self, raw_path: str, include_hidden: bool = False, preview_rows: int = 100) -> WorkbookAnalysisDto:
        started = time.perf_counter()
        path, digest = validate_workbook_path(raw_path, self.config)
        secured = time.perf_counter()
        # read_only avoids openpyxl's drawing/image loader. The service never inspects media entries.
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        parsed = time.perf_counter()
        sheets: list[SheetAnalysisDto] = []
        warnings: list[str] = []
        try:
            for worksheet in workbook.worksheets:
                if worksheet.sheet_state != "visible" and not include_hidden:
                    continue
                if worksheet.max_row is None or worksheet.max_column is None:
                    worksheet.calculate_dimension(force=True)
                max_row = worksheet.max_row or 0
                max_column = worksheet.max_column or 0
                if (max_row > self.config.max_rows or max_column > self.config.max_columns
                        or max_row * max_column > self.config.max_cells):
                    warnings.append(f"Bỏ qua sheet {worksheet.title}: vượt giới hạn hàng/cột.")
                    continue
                matrix = [list(row) for row in worksheet.iter_rows(values_only=True)]
                matrix, row_count, column_count = _trim_matrix(matrix)
                if not matrix:
                    sheets.append(SheetAnalysisDto(name=worksheet.title, state=worksheet.sheet_state, row_count=0,
                        column_count=0, merged_ranges=[], regions=[], warnings=["Sheet trống."]))
                    continue
                bold_rows = {
                    cell.row - 1
                    for row in worksheet.iter_rows(min_row=1, max_row=min(max_row, 10))
                    for cell in row
                    if cell.value is not None and cell.font and cell.font.bold
                }
                merges = _read_merge_ranges(path, worksheet._worksheet_path)
                regions = _statistical_regions(worksheet.title, matrix, merges, preview_rows)
                if not regions:
                    # Fallback cho bảng dữ liệu thông thường không có dòng mã (1), (2)...
                    for band_start, band_end in find_row_bands(matrix):
                        populated_columns = [
                            col for col in range(column_count)
                            if any(col < len(matrix[row]) and matrix[row][col] not in (None, "") for row in range(band_start, band_end + 1))
                        ]
                        if len(populated_columns) < 2:
                            continue
                        start_col, end_col = min(populated_columns), max(populated_columns)
                        header_start, header_end, confidence = choose_header_rows(matrix, band_start, band_end, bold_rows)
                        headers = flatten_headers(matrix, header_start, header_end, start_col, end_col)
                        data_start = min(header_end + 1, band_end + 1)
                        rows = [
                            [matrix[r][c] if c < len(matrix[r]) else None for c in range(start_col, end_col + 1)]
                            for r in range(data_start, min(band_end + 1, data_start + preview_rows))
                        ]
                        regions.append(TableRegionDto(sheet=worksheet.title, start_row=band_start + 1,
                            end_row=band_end + 1, start_column=start_col + 1, end_column=end_col + 1,
                            header_start_row=header_start + 1, header_end_row=header_end + 1,
                            data_start_row=data_start + 1, confidence=confidence, headers=headers,
                            source_columns=list(range(start_col + 1, end_col + 2)),
                            header_paths=[[header] for header in headers], rows=rows))
                sheets.append(SheetAnalysisDto(name=worksheet.title, state=worksheet.sheet_state,
                    row_count=row_count, column_count=column_count, merged_ranges=merges,
                    regions=regions, warnings=[] if regions else ["Không phát hiện được vùng bảng đủ tin cậy."]))
        finally:
            workbook.close()
        finished = time.perf_counter()
        return WorkbookAnalysisDto(file_name=Path(path).name, sha256=digest, sheets=sheets, warnings=warnings,
            timings_ms={"security": round((secured - started) * 1000, 2),
                        "parse": round((parsed - secured) * 1000, 2),
                        "table_detection": round((finished - parsed) * 1000, 2)},
            elapsed_ms=round((finished - started) * 1000, 2))
